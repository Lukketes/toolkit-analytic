"""
TOOLKIT PANDAS — análise exploratória feita para ler no terminal
================================================================

RODAR (relatório completo no terminal)
    python toolkit_pandas.py vendas.xlsx                    # tudo automático
    python toolkit_pandas.py base.csv -v Valor -c Tipo -d Data
    python toolkit_pandas.py base.csv -f -e Cliente         # + antifraude
    python toolkit_pandas.py base.csv -q "Valor > 1000"     # filtra antes
    python toolkit_pandas.py --demo                         # dados de exemplo
    python toolkit_pandas.py -h                             # todas as opções

EXPLORAR DEPOIS (df e rel ficam na memória)
    python -i toolkit_pandas.py base.csv
    >>> ajuda()                                  # lista de funções
    >>> mostrar(agrupar(df, "Tipo", "Valor"))
    >>> mostrar(rel.alertas)                     # com -f

USAR NO SEU PROJETO / NOTEBOOK
    from toolkit_pandas import carregar, agrupar, mostrar
    df = carregar("transacoes.csv")
    mostrar(agrupar(df, "Tipo", "Valor"))

Fluxo do relatório: visão geral -> qualidade -> ficha da variável -> outliers
-> maiores/menores -> categorias -> entidades -> tempo -> correlação
-> antifraude -> principais achados
"""
import argparse
import codecs
import csv
import difflib
import io
import json
import math
import os
import re
import shutil
import sys
import textwrap
import traceback
import unicodedata
import warnings
from datetime import date, datetime, timedelta
from pathlib import Path

import numpy as np
import pandas as pd

__version__ = "2.0"

# ---------------------------------------------------------------------------
# CONFIG — padrões usados quando a linha de comando não diz nada
# ---------------------------------------------------------------------------
BASE = Path(__file__).resolve().parent    # pasta onde este script está

ARQUIVO = None          # ex: BASE / "dados" / "transacoes.csv"  (None = informar na linha de comando)
COL_VALOR = None        # coluna numérica principal (None = escolhe sozinho)
COL_CATEGORIA = None    # ex: "Tipo"     (None = só mostra as candidatas)
COL_DATA = None         # ex: "Data"     (None = usa a única coluna de data, se houver só uma)
COL_ENTIDADE = None     # ex: "Cliente"  (quem transaciona: cliente, fornecedor, conta...)
TOP = 5                 # linhas nas listas de maiores/menores
FRAUDE = False          # True = sempre inclui a seção antifraude
LIMITE_ALCADA = None    # ex: 10000 — limite que alguém poderia querer "driblar"

# Vazio = autodetecção. Só preencha se a detecção errar, ex:
# LEITURA_KWARGS = {"sep": ";", "encoding": "latin-1"}
LEITURA_KWARGS = {}


class ErroToolkit(Exception):
    """Erro com mensagem pronta para o usuário (sem traceback)."""


class ColunaNaoEncontrada(KeyError):
    """Coluna inexistente. É um KeyError, mas com mensagem legível."""

    def __str__(self):
        return str(self.args[0]) if self.args else ""


# ===========================================================================
# 0. SAÍDA NO TERMINAL — cores, largura, linhas de status
# ===========================================================================
_SAIDA = {"cor": False, "largura": 100}

_ESTILOS = {
    "titulo": "1;36", "sub": "1", "negrito": "1", "fraco": "2", "ok": "32",
    "aviso": "33", "alerta": "1;31", "barra": "36", "barra_neg": "31",
}
_RE_ANSI = re.compile(r"\x1b\[[0-9;]*m")


def pintar(texto, estilo):
    """Aplica cor ANSI quando a saída suporta; senão devolve o texto puro."""
    if not _SAIDA["cor"] or not texto:
        return texto
    return f"\x1b[{_ESTILOS[estilo]}m{texto}\x1b[0m"


def _vlen(texto):
    """Largura visível do texto (ignora códigos de cor)."""
    return len(_RE_ANSI.sub("", texto))


def _vt_windows():
    """Liga o suporte a cores ANSI no console do Windows."""
    try:
        import ctypes
        k32 = ctypes.windll.kernel32
        handle = k32.GetStdHandle(-11)
        modo = ctypes.c_uint32()
        if not k32.GetConsoleMode(handle, ctypes.byref(modo)):
            return False
        return bool(k32.SetConsoleMode(handle, modo.value | 0x0004))
    except Exception:
        return False


def configurar_saida(cor=None, largura=None):
    """Ajusta a saída do relatório.
    cor: True/False/None (None = automático: só no terminal, respeita NO_COLOR).
    largura: colunas de texto (None = largura do terminal, entre 70 e 120)."""
    if cor is None:
        if os.environ.get("NO_COLOR"):
            cor = False
        elif os.environ.get("FORCE_COLOR"):
            cor = True
        else:
            tty = bool(getattr(sys.stdout, "isatty", lambda: False)())
            cor = tty and (_vt_windows() if os.name == "nt" else os.environ.get("TERM") != "dumb")
    if largura is None:
        largura = max(70, min(120, shutil.get_terminal_size((100, 24)).columns - 1))
    _SAIDA["cor"], _SAIDA["largura"] = bool(cor), max(40, int(largura))


def _escrever(texto, recuo=2, marcador="", estilo=None):
    """Imprime texto quebrando linhas na largura do relatório."""
    pref = " " * recuo
    if marcador:
        primeira = pref + (pintar(marcador, estilo) if estilo else marcador) + " "
        resto = pref + " " * (_vlen(marcador) + 1)
    else:
        primeira = resto = pref
    largura = max(20, _SAIDA["largura"] - len(resto))
    linhas = textwrap.wrap(str(texto), largura, break_on_hyphens=False) or [""]
    print(primeira + linhas[0])
    for linha in linhas[1:]:
        print(resto + linha)


def nota(texto="", recuo=2):
    """Texto comum do relatório."""
    _escrever(texto, recuo)


def aviso(texto, recuo=2):
    """Linha de atenção: '! texto'."""
    _escrever(texto, recuo, "!", "aviso")


def alerta(texto, recuo=2):
    """Linha de alerta forte: '!! texto'."""
    _escrever(texto, recuo, "!!", "alerta")


def _status(nivel, texto, recuo=2):
    marcas = {"ok": ("✓", "ok"), "aviso": ("!", "aviso"), "alerta": ("!!", "alerta"),
              "info": ("·", "fraco")}
    marcador, estilo = marcas.get(nivel, ("·", "fraco"))
    _escrever(texto, recuo, marcador, estilo)


def titulo(texto):
    """Cabeçalho de seção."""
    cab = f"━━ {texto} "
    print()
    print(pintar(cab + "━" * max(3, _SAIDA["largura"] - len(cab)), "titulo"))


def subtitulo(texto):
    """Cabeçalho de subseção."""
    cab = f"── {texto} "
    print()
    print("  " + pintar(cab, "sub") + pintar("─" * max(3, _SAIDA["largura"] - len(cab) - 2), "fraco"))


def par(rotulo, valor, larg=24, recuo=4):
    """Linha 'rótulo ..... valor'."""
    print(f"{' ' * recuo}{rotulo:<{larg}}{valor}")


# ===========================================================================
# 1. FORMATAÇÃO — números, datas e textos no padrão brasileiro
# ===========================================================================
_PLURAIS = {"mês": "meses", "registro": "registros", "valor": "valores"}
_MESES = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]
_DIAS = ["seg", "ter", "qua", "qui", "sex", "sáb", "dom"]


def _na(v):
    """pd.isna que não quebra com listas e textos."""
    if v is None:
        return True
    if isinstance(v, str):
        return False
    try:
        return bool(pd.isna(v))
    except (TypeError, ValueError):
        return False


def _bool(serie):
    """True/False/vazio -> bool puro (vazio = False), sem avisos em nenhuma versão do pandas."""
    if isinstance(serie, pd.DataFrame):
        return serie.apply(_bool)
    return serie.astype(object).eq(True)


def br(x, casas=2):
    """Número no formato brasileiro: 1234567.891 -> '1.234.567,89'.
    Inteiros saem sem casas; None/NaN viram '-'."""
    if isinstance(x, str):
        return x
    if _na(x):
        return "-"
    if isinstance(x, (bool, np.bool_)):
        return "sim" if x else "não"
    if isinstance(x, (int, np.integer)):
        x, casas = int(x), 0
    else:
        x = float(x)
        if math.isinf(x):
            return "∞" if x > 0 else "-∞"
    s = f"{x:,.{casas}f}".replace(",", "\0").replace(".", ",").replace("\0", ".")
    if s.startswith("-") and not any(c in "123456789" for c in s):
        s = s[1:]                                   # nada de "-0,00"
    return s


def plural(n, singular, plural_=None):
    """'1 valor', '3 valores': plural(3, 'valor', 'valores')."""
    return f"{br(n)} {singular if n == 1 else (plural_ or singular + 's')}"


def pct(x, casas=1):
    """Percentual no formato BR: 12.345 -> '12,3%'."""
    return "-" if _na(x) else f"{br(float(x), casas)}%"


def pct_sinal(x, casas=1):
    """Variação com sinal: 12.3 -> '+12,3%'."""
    if _na(x) or (isinstance(x, float) and math.isinf(x)):
        return "-"
    return ("+" if x > 0 else "") + pct(x, casas)


def data_br(v):
    """Timestamp -> '31/01/2026' (ou '31/01/2026 14:30' se tiver hora)."""
    if _na(v):
        return "-"
    v = pd.Timestamp(v)
    if v.hour == 0 and v.minute == 0 and v.second == 0:
        return v.strftime("%d/%m/%Y")
    return v.strftime("%d/%m/%Y %H:%M")


def cortar(texto, n=30):
    """Encurta o texto para caber em n caracteres (e tira quebras de linha)."""
    texto = re.sub(r"\s+", " ", str(texto))
    return texto if len(texto) <= n else texto[: n - 1] + "…"


def _casas(serie):
    """Casas decimais adequadas à escala dos valores (0,0012 precisa de mais que 2)."""
    s = _num(serie).abs()
    s = s[s > 0]
    if s.empty:
        return 2
    m = s.median()
    return 2 if m >= 1 else 3 if m >= 0.1 else 4 if m >= 0.001 else 6


def _bytes(n):
    for unidade in ("B", "KB", "MB", "GB"):
        if n < 1024 or unidade == "GB":
            return f"{br(n, 0 if unidade == 'B' else 1)} {unidade}"
        n /= 1024


def _celula(v, casas=2, max_cel=30):
    """Formata um valor qualquer para caber numa célula de tabela."""
    if isinstance(v, str):
        return cortar(v, max_cel)
    if _na(v):
        return "-"
    if isinstance(v, (pd.Timestamp, datetime, np.datetime64)):
        return data_br(v)
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, (bool, np.bool_)):
        return "sim" if v else "não"
    if isinstance(v, (int, np.integer)):
        return br(v)
    if isinstance(v, (float, np.floating)):
        return br(v, casas)
    return cortar(v, max_cel)


def barra(valor, maximo, largura=20):
    """Barra de texto proporcional: barra(50, 100) -> '██████████'."""
    try:
        if _na(valor) or _na(maximo) or not maximo:
            return ""
        frac = min(abs(float(valor)) / abs(float(maximo)), 1.0)
    except (TypeError, ValueError):
        return ""
    meias = round(frac * largura * 2)
    if meias == 0 and valor:
        meias = 1
    return pintar("█" * (meias // 2) + ("▌" if meias % 2 else ""),
                  "barra_neg" if valor < 0 else "barra")


def _barra_ref(valor, ref, maximo, largura=20):
    """Barra com um marcador '│' na posição de referência (ex: o esperado)."""
    if not maximo:
        return ""
    cheio = int(round(min(valor / maximo, 1) * largura))
    marca = min(int(round(ref / maximo * largura)), largura)
    chars = ["█" if i < cheio else " " for i in range(max(cheio, marca + 1))]
    chars[marca] = "│"
    return pintar("".join(chars).rstrip(), "barra")


# ===========================================================================
# 2. TABELAS
# ===========================================================================
def tabela(cabecalho, linhas, alinh=None, recuo=2):
    """Imprime uma tabela alinhada a partir de listas.
    alinh: 'l' (esquerda) ou 'r' (direita) por coluna. Colunas que não cabem
    na largura do terminal são omitidas, com aviso."""
    cab = [str(c) for c in cabecalho]
    linhas = [[str(c) for c in linha] for linha in linhas]
    n = len(cab)
    alinh = list(alinh) if alinh else ["l"] * n
    larg = [max([_vlen(cab[i])] + [_vlen(lin[i]) for lin in linhas]) for i in range(n)]
    cabe = _SAIDA["largura"] - recuo
    k = n
    while k > 1 and sum(larg[:k]) + 2 * (k - 1) > cabe:
        k -= 1

    def fmt(celulas):
        partes = []
        for i in range(k):
            pad = " " * (larg[i] - _vlen(celulas[i]))
            partes.append(pad + celulas[i] if alinh[i] == "r" else celulas[i] + pad)
        return (" " * recuo + "  ".join(partes)).rstrip()

    print(pintar(fmt(cab), "negrito"))
    print(" " * recuo + pintar("  ".join("─" * w for w in larg[:k]), "fraco"))
    for linha in linhas:
        print(fmt(linha))
    if not linhas:
        nota("(nenhuma linha)", recuo)
    if k < n:
        resto = [c for c in cab[k:] if c]
        if resto:
            nomes = ", ".join(resto[:6]) + ("…" if len(resto) > 6 else "")
            nota(f"(+{len(resto)} colunas não couberam na largura: {nomes})", recuo)


def _inteiro_sem_milhar(serie):
    """Inteiros que são códigos, linhas ou anos: mostrar '2026', não '2.026'."""
    if _parece_codigo_nome(serie.name) or set(_tokens(serie.name)) & {"linha", "indice", "ano", "grupo"}:
        return True
    s = serie.dropna()
    return len(s) > 0 and bool(s.between(1900, 2100).all())


def _formatador(serie, casas=2, max_cel=30):
    if pd.api.types.is_bool_dtype(serie):
        return _celula
    if pd.api.types.is_integer_dtype(serie):
        if _inteiro_sem_milhar(serie):
            return lambda v: "-" if _na(v) else str(int(v))
        return _celula
    if pd.api.types.is_float_dtype(serie):
        return lambda v: _celula(v, casas)
    return lambda v: _celula(v, casas, max_cel)


def mostrar(obj, casas=2, indice=True, recuo=2, max_linhas=20, prioridade=None, max_cel=30):
    """Imprime DataFrame/Series alinhado, com números no formato brasileiro.
    max_linhas: limite de linhas (None = todas).
    prioridade: colunas que vêm primeiro quando nem todas cabem na largura.
    max_cel: largura máxima de cada célula de texto."""
    if isinstance(obj, pd.Series):
        obj = obj.to_frame(obj.name if obj.name is not None else "valor")
    if not isinstance(obj, pd.DataFrame):
        print(obj)
        return
    df, total = obj, len(obj)
    if max_linhas is not None and total > max_linhas:
        df = df.head(max_linhas)
    if prioridade:
        frente = list(dict.fromkeys(c for c in prioridade if c in df.columns))
        df = df[frente + [c for c in df.columns if c not in frente]]
    nomes = [" / ".join(map(str, c)) if isinstance(c, tuple) else str(c) for c in df.columns]
    fmts = [_formatador(df.iloc[:, i], casas, max_cel) for i in range(df.shape[1])]
    alinh = ["r" if pd.api.types.is_numeric_dtype(df.iloc[:, i])
             and not pd.api.types.is_bool_dtype(df.iloc[:, i]) else "l"
             for i in range(df.shape[1])]
    linhas = [[f(v) for f, v in zip(fmts, row)] for row in df.itertuples(index=False, name=None)]
    if indice:
        nome_idx = " / ".join(str(n) for n in df.index.names if n is not None)
        rot = [" / ".join(_celula(x, casas) for x in i) if isinstance(i, tuple)
               else (str(i) if isinstance(i, (int, np.integer)) else _celula(i, casas))
               for i in df.index]
        linhas = [[r] + lin for r, lin in zip(rot, linhas)]
        nomes = [nome_idx] + nomes
        alinh = ["l"] + alinh
    tabela(nomes, linhas, alinh, recuo)
    if len(df) < total:
        nota(f"(mostrando {len(df)} de {br(total)} linhas; mostrar(x, max_linhas=None) mostra tudo)",
             recuo)


# ===========================================================================
# 3. COLUNAS — nomes, papéis e candidatas
# ===========================================================================
def _sem_acento(texto):
    texto = unicodedata.normalize("NFKD", str(texto))
    return "".join(c for c in texto if not unicodedata.combining(c))


def _normalizar(texto):
    """'  Valor Total ' -> 'valortotal' (sem acento, espaço ou maiúscula)."""
    return re.sub(r"[^a-z0-9]", "", _sem_acento(texto).lower())


def _tokens(nome):
    """Palavras do nome da coluna: 'idCliente' -> ['id', 'cliente'],
    'VLR_TOTAL' -> ['vlr', 'total'], 'CPFCliente' -> ['cpf', 'cliente']."""
    s = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", str(nome))
    s = re.sub(r"([A-Z]+)([A-Z][a-z])", r"\1 \2", s)
    return [t for t in re.split(r"[^a-z0-9]+", _sem_acento(s).lower()) if t]


def _eh_texto(serie):
    return (pd.api.types.is_object_dtype(serie) or pd.api.types.is_string_dtype(serie)) \
        and not isinstance(serie.dtype, pd.CategoricalDtype)


def _eh_numerica(serie):
    return pd.api.types.is_numeric_dtype(serie) and not pd.api.types.is_bool_dtype(serie)


def _num(serie):
    """Série como float64 (NaN no lugar de vazios/textos), segura para numpy."""
    s = pd.to_numeric(serie, errors="coerce")
    return pd.Series(s.to_numpy(dtype="float64", na_value=np.nan), index=serie.index,
                     name=serie.name)


def _nunique(serie):
    try:
        return int(serie.nunique(dropna=True))
    except TypeError:                                 # listas/dicts dentro das células
        return int(serie.dropna().astype(str).nunique())


def resolver_coluna(df, nome):
    """Acha a coluna mesmo com diferença de maiúscula, acento ou espaço.
    Se não achar, levanta erro sugerindo nomes parecidos."""
    if nome in df.columns:
        return nome
    alvo = _normalizar(nome)
    achadas = [c for c in df.columns if _normalizar(c) == alvo]
    if len(achadas) == 1:
        return achadas[0]
    normais = {_normalizar(c): c for c in df.columns}
    parecidas = difflib.get_close_matches(alvo, list(normais), n=3, cutoff=0.6)
    parecidas += [n for n in normais if alvo and alvo in n and n not in parecidas]
    msg = f"Coluna '{nome}' não existe."
    if parecidas:
        msg += " Você quis dizer: " + ", ".join(repr(normais[p]) for p in parecidas[:3]) + "?"
    msg += "\nColunas disponíveis: " + ", ".join(map(str, df.columns))
    raise ColunaNaoEncontrada(msg)


_ID_TOKENS = {"id", "cod", "codigo", "code", "uuid", "guid", "chave", "key", "pk",
              "matricula", "nsu", "protocolo", "cpf", "cnpj", "cep", "telefone", "celular",
              "fone", "rg", "pis", "nis"}
_NUM_ID_TOKENS = {"num", "numero", "nro", "no", "n"}
_NUM_ID_OBJETOS = {"pedido", "nota", "nf", "nfe", "documento", "doc", "contrato", "processo",
                   "conta", "serie", "protocolo", "transacao", "registro", "cliente", "ordem",
                   "boleto", "titulo", "fatura", "cartao"}
# colunas que são códigos mesmo parecendo número: nunca converter para número
_CODIGO_TOKENS = {"cpf", "cnpj", "cep", "telefone", "tel", "celular", "fone", "whatsapp", "rg",
                  "pis", "nis", "matricula", "agencia", "conta", "cartao", "nsu", "protocolo",
                  "chave", "barras", "ean", "gtin", "ncm", "cfop", "renavam", "placa"}


def _parece_codigo_nome(nome):
    """O NOME indica identificador/código? ('ID_Cliente', 'clienteId', 'Nº Pedido')."""
    toks = _tokens(nome)
    if not toks:
        return False
    if toks[0] in _ID_TOKENS or toks[-1] in _ID_TOKENS:
        return True
    return len(toks) >= 2 and toks[0] in _NUM_ID_TOKENS and toks[1] in _NUM_ID_OBJETOS


def parece_id(df, col):
    """Coluna que é só identificador: nome de ID/código ('ID_Cliente', 'cod_produto'),
    inteiros únicos em sequência, ou códigos de texto únicos sem espaço ('TX000123').
    'Idade', 'Pago' e 'Valid' NÃO são ID."""
    if _parece_codigo_nome(col):
        return True
    if set(_tokens(col)) & (_PISTAS_VALOR | _PISTAS_CONTAGEM):
        return False                       # 'Idade', 'Qtd', 'Valor': medida, mesmo se ordenada
    s = df[col].dropna()
    if len(s) < 10:
        return False
    try:
        if not s.is_unique:
            return False
    except TypeError:
        return False
    if _eh_numerica(s):
        s = _num(s)
        if not (s % 1 == 0).all():
            return False
        if s.is_monotonic_increasing or s.is_monotonic_decreasing:
            return True
        digitos = np.floor(np.log10(s.abs().clip(lower=1)))
        return bool(s.min() >= 1000 and digitos.min() == digitos.max())
    if _eh_texto(s):
        t = (s.sample(20_000, random_state=0) if len(s) > 20_000 else s).astype(str)
        return bool(t.str.len().median() <= 40 and (~t.str.contains(" ", regex=False)).mean() > 0.95)
    return False


_PISTAS_VALOR = {"valor", "vlr", "vl", "total", "preco", "price", "amount", "value", "montante",
                 "quantia", "faturamento", "receita", "saldo", "custo", "venda", "vendas",
                 "pagamento", "pago", "liquido", "bruto", "debito", "credito", "despesa"}
_PISTAS_CONTAGEM = {"qtd", "qtde", "quantidade", "ano", "mes", "dia", "idade", "parcela",
                    "parcelas", "score", "itens"}


def escolher_coluna_valor(df):
    """Coluna numérica mais provável de ser 'o valor': pelo nome (valor, total,
    preço, vlr...), depois preferindo decimais com muitos valores distintos."""
    cands = [c for c in df.columns if _eh_numerica(df[c]) and not parece_id(df, c)]
    if not cands:
        return None

    def pontos(c):
        toks, s = set(_tokens(c)), df[c]
        p = 10 if toks & _PISTAS_VALOR else 0
        p += 2 if toks & {"valor", "vlr", "vl"} else 0
        p += 2 if pd.api.types.is_float_dtype(s) else 0
        p -= 3 if toks & _PISTAS_CONTAGEM else 0
        return p + min(_nunique(s) / max(len(s), 1), 1)

    return max(cands, key=pontos)


_ENTIDADE_TOKENS = {"cliente", "fornecedor", "favorecido", "beneficiario", "pagador", "recebedor",
                    "conta", "usuario", "vendedor", "funcionario", "empresa", "parceiro", "titular",
                    "comprador", "colaborador", "credor", "devedor", "sacado", "cedente", "emitente",
                    "destinatario", "remetente", "prestador", "tomador", "correntista", "loja"}
_DATA_NAO_EVENTO = {"nascimento", "nasc", "vencimento", "venc", "validade", "expiracao",
                    "admissao", "demissao", "previsao", "prazo", "cadastro", "abertura"}


def colunas_data(df):
    """Colunas que são (ou parecem) datas."""
    out = []
    for col in df.columns:
        s = df[col]
        if pd.api.types.is_datetime64_any_dtype(s):
            out.append(col)
        elif _eh_texto(s):
            textos, pesos, _ = _amostra_pesada(s, 500)
            if _familia_data(textos, pesos)[1] >= 0.9:
                out.append(col)
    return out


def colunas_categoricas(df, max_unicos=30):
    """Colunas de texto/booleanas com poucos valores distintos (bons grupos)."""
    out = []
    for col in df.columns:
        s = df[col]
        if pd.api.types.is_datetime64_any_dtype(s):
            continue
        if pd.api.types.is_bool_dtype(s) or isinstance(s.dtype, pd.CategoricalDtype) or _eh_texto(s):
            n, preenchidos = _nunique(s), int(s.notna().sum())
            if 2 <= n <= max_unicos and (n <= preenchidos * 0.5 or preenchidos < 20) \
                    and not parece_id(df, col):
                out.append(col)
    return [c for c in out if c not in set(colunas_data(df[out]))] if out else out


def colunas_entidade(df, pii=None):
    """Colunas que identificam QUEM transaciona (cliente, fornecedor, conta, CPF...)."""
    pii = detectar_pii(df) if pii is None else pii
    out = []
    for col in df.columns:
        s = df[col]
        if pd.api.types.is_datetime64_any_dtype(s) or (_eh_numerica(s) and not _parece_codigo_nome(col)):
            continue
        n, preenchidos = _nunique(s), int(s.notna().sum())
        if n < 10 or n >= preenchidos:
            continue
        if set(_tokens(col)) & _ENTIDADE_TOKENS or pii.get(col) in ("CPF", "CNPJ"):
            out.append(col)
    return out


def _nome_tipo(serie):
    if pd.api.types.is_bool_dtype(serie):
        return "booleano"
    if isinstance(serie.dtype, pd.CategoricalDtype):
        return "categoria"
    if pd.api.types.is_integer_dtype(serie):
        return "inteiro"
    if pd.api.types.is_float_dtype(serie):
        return "decimal"
    if pd.api.types.is_datetime64_any_dtype(serie):
        return "data"
    if pd.api.types.is_timedelta64_dtype(serie):
        return "duração"
    if pd.api.types.is_object_dtype(serie):
        amostra = serie.dropna().head(500)
        if len(amostra) and not amostra.map(lambda v: isinstance(v, str)).all():
            return "misto"
    return "texto"


_VALORES_BINARIOS = {"sim", "nao", "s", "n", "true", "false", "verdadeiro", "falso", "0", "1",
                     "y", "yes", "no", "ativo", "inativo", "x", ""}


def _papel(df, col, pii):
    """Que dado é esse? vazia, constante, data, pessoal, binária, identificador,
    numérica, categórica ou texto livre."""
    s = df[col]
    preenchidos = s.dropna()
    if preenchidos.empty:
        return "vazia"
    n = _nunique(s)
    if n == 1 and len(preenchidos) > 1:
        return "constante"
    if pd.api.types.is_datetime64_any_dtype(s):
        return "data"
    if col in pii:
        return f"pessoal: {pii[col]}"
    if pd.api.types.is_bool_dtype(s) or (n == 2 and {_normalizar(v) for v in preenchidos.unique()}
                                         <= _VALORES_BINARIOS):
        return "binária"
    if parece_id(df, col):
        return "identificador"
    if _eh_numerica(s):
        return "numérica"
    if set(_tokens(col)) & _ENTIDADE_TOKENS and n >= 10:
        return "entidade"
    if n <= 30 or n <= 0.5 * len(preenchidos):
        return "categórica"
    return "texto livre"


# ===========================================================================
# 4. CARGA — ler qualquer arquivo e arrumar números BR e datas
# ===========================================================================
_EXT_CSV = {".csv", ".txt", ".tsv"}
_EXT_EXCEL = {".xlsx", ".xlsm", ".xls", ".xlsb", ".ods"}
_EXT_JSON = {".json", ".jsonl", ".ndjson"}
_EXT_PARQUET = {".parquet", ".pq"}
_EXTENSOES = _EXT_CSV | _EXT_EXCEL | _EXT_JSON | _EXT_PARQUET
_DEPENDENCIAS = {".xls": "xlrd", ".xlsb": "pyxlsb", ".ods": "odfpy", ".xlsx": "openpyxl",
                 ".xlsm": "openpyxl", ".parquet": "pyarrow", ".pq": "pyarrow"}


def _detectar_csv(caminho, max_bytes=512_000):
    """Descobre encoding, separador e a linha do cabeçalho olhando o começo do arquivo.
    Devolve (sep, encoding, indice_da_linha_do_cabecalho)."""
    with open(caminho, "rb") as f:
        bruto = f.read(max_bytes)
    truncado = len(bruto) == max_bytes
    encoding, texto = "latin-1", bruto.decode("latin-1")
    for enc in ("utf-8-sig", "cp1252"):
        try:
            # incremental: tolera um caractere cortado no fim da amostra
            texto = codecs.getincrementaldecoder(enc)().decode(bruto, final=not truncado)
            encoding = enc
            break
        except UnicodeDecodeError:
            continue
    linhas = texto.splitlines()
    if truncado and linhas:
        linhas = linhas[:-1]
    linhas = linhas[:60]

    melhor = None
    for sep in (";", ",", "\t", "|"):
        try:
            contagens = [len(c) for c in csv.reader(linhas, delimiter=sep)]
        except csv.Error:
            continue
        uteis = [n for n in contagens if n > 1]
        preenchidas = [n for n in contagens if n > 0]
        if not uteis:
            continue
        moda = max(set(uteis), key=lambda n: (uteis.count(n), n))
        chave = (round(contagens.count(moda) / len(preenchidas), 2), moda)
        if melhor is None or chave > melhor[0]:
            melhor = (chave, sep, moda)
    if melhor is None:
        return ",", encoding, 0
    _, sep, moda = melhor
    contagens = [len(c) for c in csv.reader(linhas, delimiter=sep)]
    cabecalho = next((i for i, n in enumerate(contagens) if n == moda), 0)
    return sep, encoding, cabecalho


def _ler_csv(caminho, info, **kwargs):
    sep, enc, cab = _detectar_csv(caminho)
    enc_do_usuario = "encoding" in kwargs
    kwargs.setdefault("sep", sep)
    kwargs.setdefault("encoding", enc)
    kwargs.setdefault("dtype", str)                # números/datas são convertidos depois, com regras BR
    if cab and "skiprows" not in kwargs and "header" not in kwargs:
        kwargs["skiprows"] = cab
    tentativas = [kwargs["encoding"]]
    if not enc_do_usuario:
        tentativas += [e for e in ("cp1252", "latin-1") if e != kwargs["encoding"]]
    df = None
    for enc in tentativas:
        kwargs["encoding"] = enc
        try:
            df = pd.read_csv(caminho, **kwargs)
            break
        except UnicodeDecodeError:
            continue
        except pd.errors.ParserError:
            ruins = []
            kw = dict(kwargs, engine="python", on_bad_lines=lambda linha: ruins.append(linha))
            kw.pop("low_memory", None)
            df = pd.read_csv(caminho, **kw)
            info["mensagens"].append((
                "aviso", f"{br(len(ruins))} linhas mal formadas (com colunas a mais) foram "
                         "ignoradas; os números de linha podem ficar deslocados."))
            break
    if df is None:
        raise ErroToolkit(f"Não consegui decodificar o texto de {Path(caminho).name}. "
                          "Tente --encoding utf-8 ou --encoding latin-1.")
    skip = kwargs.get("skiprows") if isinstance(kwargs.get("skiprows"), int) else 0
    info.update(formato="CSV", sep=kwargs["sep"], encoding=kwargs["encoding"],
                linha_cabecalho=skip + 1, linha_offset=skip + 2)
    return df


def _linha_cabecalho(bruto, max_busca=30):
    """Em planilhas com título em cima, acha a linha que é o cabeçalho de verdade:
    a primeira linha 'cheia' composta quase só de textos."""
    if bruto.empty:
        return 0
    cheios = bruto.notna().sum(axis=1).to_numpy()
    alvo = cheios.max()
    if alvo < 2:
        return 0
    for i in range(min(len(bruto), max_busca)):
        if cheios[i] < max(2, 0.6 * alvo):
            continue
        valores = [v for v in bruto.iloc[i].tolist() if not _na(v)]
        textos = sum(isinstance(v, str) and not _RE_NUMERO_SIMPLES.fullmatch(v.strip())
                     for v in valores)
        return i if textos >= 0.8 * len(valores) else 0
    return 0


def _resolver_aba(abas, aba):
    if aba is None:
        return abas[0]
    if isinstance(aba, int) or str(aba).strip().isdigit():
        i = int(aba)
        if 1 <= i <= len(abas):
            return abas[i - 1]
        raise ErroToolkit(f"A planilha tem {len(abas)} abas; --aba {i} não existe. "
                          f"Abas: {', '.join(abas)}")
    for a in abas:
        if a == aba or _normalizar(a) == _normalizar(aba):
            return a
    raise ErroToolkit(f"Aba '{aba}' não existe. Abas: {', '.join(abas)}")


def _ler_excel(caminho, info, aba=None, cabecalho=None, **kwargs):
    with pd.ExcelFile(caminho) as xl:
        abas = list(xl.sheet_names)
        nome = _resolver_aba(abas, aba)
        if cabecalho:
            h = int(cabecalho) - 1
        elif "header" in kwargs:
            h = kwargs.pop("header") or 0
        else:
            h = _linha_cabecalho(xl.parse(nome, header=None, nrows=30))
        df = xl.parse(nome, header=h, **kwargs)
    info.update(formato="Excel", aba=nome, abas=abas, linha_cabecalho=h + 1, linha_offset=h + 2)
    return df


def _achar_registros(dados):
    """Num JSON, acha a lista de registros (mesmo dentro de {'data': [...]})."""
    if isinstance(dados, list):
        return dados if all(isinstance(d, dict) for d in dados[:100]) else None
    if isinstance(dados, dict):
        listas = [v for v in dados.values() if isinstance(v, list) and v and isinstance(v[0], dict)]
        if listas:
            return max(listas, key=len)
    return None


def _ler_json(caminho, info, **kwargs):
    info["formato"] = "JSON"
    if Path(caminho).suffix.lower() in (".jsonl", ".ndjson"):
        return pd.read_json(caminho, lines=True, **kwargs)
    try:
        with open(caminho, encoding="utf-8-sig") as f:
            dados = json.load(f)
    except json.JSONDecodeError:
        return pd.read_json(caminho, lines=True, **kwargs)     # JSON Lines com extensão .json
    registros = _achar_registros(dados)
    if registros is not None:
        return pd.json_normalize(registros)                   # aninhados viram 'cliente.nome'
    return pd.DataFrame(dados)


def _carregar(caminho, aba=None, cabecalho=None, **kwargs):
    """Lê o arquivo e devolve (df, info). info guarda formato, separador,
    encoding, aba, linha do cabeçalho e as mensagens da limpeza."""
    caminho = Path(caminho).expanduser()
    if not caminho.exists():
        raise FileNotFoundError(_msg_arquivo_inexistente(caminho))
    if caminho.is_dir():
        raise ErroToolkit(f"{caminho} é uma pasta. Aponte para o arquivo."
                          + _sugerir_arquivos(caminho))
    ext = caminho.suffix.lower()
    if ext not in _EXTENSOES:
        raise ErroToolkit(f"Extensão não suportada: '{ext}' ({caminho.name}).\n"
                          f"Suportadas: {', '.join(sorted(_EXTENSOES))}")
    info = {"arquivo": caminho.name, "caminho": str(caminho), "mensagens": []}
    try:
        if ext in _EXT_CSV:
            df = _ler_csv(caminho, info, **kwargs)
        elif ext in _EXT_EXCEL:
            df = _ler_excel(caminho, info, aba=aba, cabecalho=cabecalho, **kwargs)
        elif ext in _EXT_JSON:
            df = _ler_json(caminho, info, **kwargs)
        else:
            info["formato"] = "Parquet"
            df = pd.read_parquet(caminho, **kwargs)
    except ImportError as erro:
        pacote = _DEPENDENCIAS.get(ext, "")
        raise ErroToolkit(f"Falta uma biblioteca para ler '{ext}'. Instale com:\n"
                          f"    pip install {pacote}\n({erro})") from erro
    except PermissionError as erro:
        raise ErroToolkit(f"Sem permissão para ler {caminho.name}. Se ele estiver aberto "
                          "no Excel, feche e rode de novo.") from erro
    except pd.errors.EmptyDataError as erro:
        raise ErroToolkit(f"{caminho.name} está vazio.") from erro
    df = preparar(df, verbose=False, preferir_br=info.get("sep") != ",", _info=info)
    return df, info


def carregar(caminho, verbose=True, aba=None, cabecalho=None, **kwargs):
    """Lê .csv/.txt/.tsv, .xlsx/.xls/.xlsm, .json/.jsonl ou .parquet pela extensão.

    - CSV: detecta separador, encoding e linhas de título antes do cabeçalho.
    - Excel: aba = nome ou número (1 = primeira); acha o cabeçalho mesmo com
      título em cima. cabecalho = número da linha do cabeçalho (1 = primeira).
    - Todos: 'R$ 1.234,56' vira número, '31/01/2026' vira data, IDs com zero à
      esquerda ('00123') continuam texto. kwargs vão para o leitor do pandas.
    """
    df, info = _carregar(caminho, aba=aba, cabecalho=cabecalho, **kwargs)
    if verbose:
        _imprimir_carga(info, df)
    return df


# --- números em texto -------------------------------------------------------
_RE_NUMERO_SIMPLES = re.compile(r"[-+]?\d+([.,]\d+)?")
_NUCLEO_BR = r"(?:\d{1,3}(?:\.\d{3})+|\d+)(?:,\d+)?"
_NUCLEO_US = r"(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?"


def _padrao_numero(nucleo):
    corpo = rf"[-+]?\s*(?:R\$|US\$|\$)?\s*[-+]?\s*{nucleo}\s*%?"
    return rf"(?:{corpo}|\(\s*{corpo}\s*\)|{corpo}\s*-)"


_PADRAO_BR = _padrao_numero(_NUCLEO_BR)
_PADRAO_US = _padrao_numero(_NUCLEO_US)
_MARCADORES_VAZIO = {"", "-", "--", "—", "n/a", "na", "n/d", "nd", "null", "none", "nan", "nat",
                     "s/n", "s/d", "sem valor"}
_ERROS_EXCEL = {"#n/d", "#n/a", "#valor!", "#value!", "#ref!", "#div/0!", "#nome?", "#name?",
                "#num!", "#nulo!", "#null!"}


def _so_textos(serie):
    """Série object só com os valores que são str (o resto vira NaN)."""
    s = serie.astype(object)
    if pd.api.types.infer_dtype(s, skipna=True) in ("string", "empty"):
        return s                                       # caminho rápido: já é tudo texto
    return s.where(s.map(lambda v: isinstance(v, str)))


def _por_unicos(serie, funcao):
    """Aplica `funcao` (Series -> Series) só nos valores DISTINTOS e espalha o
    resultado de volta. 1 milhão de linhas com 10 valores distintos = 10 cálculos."""
    try:
        codigos, unicos = pd.factorize(serie)
    except TypeError:                                  # células com lista/dict
        return funcao(serie)
    res = funcao(pd.Series(np.asarray(unicos, dtype=object))).reset_index(drop=True)
    out = res.reindex(codigos)
    out.index, out.name = serie.index, serie.name
    return out


def _amostra_texto(serie, n=2000, completo=False):
    """Até n valores de texto distintos (sorteados, se houver mais), sem espaços
    nas pontas e sem marcadores de vazio. Em bases grandes olha 200 mil linhas
    sorteadas, a não ser que completo=True."""
    if not completo and len(serie) > 200_000:
        serie = serie.iloc[np.random.default_rng(0).integers(0, len(serie), 200_000)]
    try:
        unicos = pd.unique(serie.dropna().to_numpy(dtype=object))
    except TypeError:
        unicos = serie.dropna().astype(str).unique()
    if len(unicos) > 5 * n:
        unicos = unicos[np.random.default_rng(0).choice(len(unicos), 5 * n, replace=False)]
    vistos, out = set(), []
    for v in unicos:
        if isinstance(v, str):
            t = v.strip()
            if t.lower() not in _MARCADORES_VAZIO and t.lower() not in _ERROS_EXCEL and t not in vistos:
                vistos.add(t)
                out.append(t)
                if len(out) >= n:
                    break
    return pd.Series(out, dtype=object)


def _amostra_pesada(serie, n_max=20_000):
    """(textos distintos, pesos, completo): cada texto com o nº de linhas em que
    aparece, para que '95% dos valores' signifique 95% das LINHAS. Em bases grandes,
    conta numa amostra de 200 mil linhas sorteadas. completo = todos os distintos."""
    completo = len(serie) <= 200_000
    if not completo:
        serie = serie.iloc[np.random.default_rng(0).integers(0, len(serie), 200_000)]
    vc = _so_textos(serie).value_counts()
    if vc.empty:
        return pd.Series([], dtype=object), np.array([], dtype=float), completo
    textos = pd.Series(vc.index.to_numpy(dtype=object)).str.strip()
    g = pd.Series(vc.to_numpy(dtype=float)).groupby(textos.to_numpy()).sum()
    minus = pd.Series(g.index.astype(str)).str.lower().to_numpy()
    g = g[~np.isin(minus, list(_MARCADORES_VAZIO | _ERROS_EXCEL))]
    if len(g) > n_max:
        g, completo = g.sample(n_max, random_state=0), False
    return pd.Series(g.index.to_numpy(dtype=object)), g.to_numpy(dtype=float), completo


def _falhas(original, convertido):
    """Valores preenchidos que não converteram (ignorando '-', 'n/d', '#N/D'...).
    Devolve (textos que falharam, quantidade de erros do Excel)."""
    cand = original[original.notna() & convertido.isna()]
    if cand.empty:
        return cand.astype(object), 0
    t = cand.astype(str).str.strip()
    minus = t.str.lower()
    erro_excel = minus.isin(_ERROS_EXCEL)
    return t[~minus.isin(_MARCADORES_VAZIO) & ~erro_excel], int(erro_excel.sum())


def _formato_numero(textos, preferir_br=True, pesos=None):
    """Decide se os textos são números BR (1.234,56), US (1,234.56) ou nada.
    Exige 95% no mesmo formato: dos valores distintos, ou das linhas (pesos)
    quando os que não batem são no máximo 5 valores distintos."""
    if textos.empty:
        return None
    w = np.ones(len(textos)) if pesos is None else np.asarray(pesos, dtype=float)
    ok_br = textos.str.fullmatch(_PADRAO_BR).to_numpy(dtype=bool)
    ok_us = textos.str.fullmatch(_PADRAO_US).to_numpy(dtype=bool)
    so_br, so_us = int((ok_br & ~ok_us).sum()), int((ok_us & ~ok_br).sum())
    if so_br != so_us:
        formato = "br" if so_br > so_us else "us"
    else:
        formato = "br" if preferir_br else "us"
    ok = ok_br if formato == "br" else ok_us
    por_linha = w[ok].sum() / w.sum() >= 0.95 and int((~ok).sum()) <= 5
    return formato if ok.mean() >= 0.95 or por_linha else None


def converter_numero(serie, formato="br"):
    """Texto -> número. formato 'br' ('R$ 1.234,56', '(500,00)', '12,5%', '1.234,56-')
    ou 'us' ('1,234.56'). O que não for número vira NaN. Números já numéricos ficam."""
    if _eh_numerica(serie):
        return serie
    padrao = _PADRAO_BR if formato == "br" else _PADRAO_US

    def converter_unicos(u):
        eh_num = u.map(lambda v: isinstance(v, (int, float, np.integer, np.floating))
                       and not isinstance(v, (bool, np.bool_))).astype(bool)
        s = _so_textos(u).str.replace("\u00a0", " ", regex=False).str.strip()
        ok = _bool(s.str.fullmatch(padrao))
        negativo = ((s.str.startswith("(") & s.str.endswith(")")) | s.str.endswith("-"))
        negativo = _bool(negativo)
        limpo = s.str.replace(r"R\$|US\$|\$|%|\s|\(|\)", "", regex=True).str.replace(r"-$", "", regex=True)
        if formato == "br":
            limpo = limpo.str.replace(".", "", regex=False).str.replace(",", ".", regex=False)
        else:
            limpo = limpo.str.replace(",", "", regex=False)
        num = pd.to_numeric(limpo.where(ok), errors="coerce").astype("float64")
        num = num.where(~negativo, -num.abs())
        return num.where(~eh_num, pd.to_numeric(u.where(eh_num), errors="coerce"))

    num = _por_unicos(serie.astype(object), converter_unicos)
    return pd.Series(num.to_numpy(dtype="float64", na_value=np.nan), index=serie.index,
                     name=serie.name)


def _converter_numeros(df, msgs, preferir_br=True):
    for col in df.columns:
        s = df[col]
        if not _eh_texto(s) or set(_tokens(col)) & _CODIGO_TOKENS:
            continue
        textos, pesos, completo = _amostra_pesada(s)
        if textos.empty:
            continue
        formato = _formato_numero(textos, preferir_br, pesos)
        if formato is None:
            continue
        # códigos com zero à esquerda ou longos demais (cartão, código de barras) ficam texto
        distintos = textos if completo else _amostra_texto(s, 10 ** 9, completo=True)
        if distintos.str.fullmatch(r"[-+]?0\d+").any() or distintos.str.fullmatch(r"\d{15,}").any():
            continue
        num = converter_numero(s, formato)
        falhou, erros_excel = _falhas(s, num)
        tem_decimal = distintos.str.contains("," if formato == "br" else ".", regex=False).any()
        if not tem_decimal and num.notna().all() and (num % 1 == 0).all():
            num = num.astype("int64")
        df[col] = num
        exemplo = textos.iloc[int(np.argmax(pesos))]
        if not re.fullmatch(r"[-+]?\d+", exemplo):
            msgs.append(("ok", f"{col}: texto '{cortar(exemplo, 18)}' convertido para número."))
        if len(falhou):
            ex = ", ".join(repr(cortar(v, 15)) for v in falhou.drop_duplicates().head(3))
            msgs.append(("aviso", f"{col}: {plural(len(falhou), 'valor não numérico virou', 'valores não numéricos viraram')} "
                                  f"vazio (ex: {ex})."))
        if erros_excel:
            msgs.append(("aviso", f"{col}: {br(erros_excel)} células com erro de fórmula do "
                                  "Excel (#N/D, #VALOR!...) viraram vazio."))


# --- datas em texto -----------------------------------------------------------
_RE_PARECE_DATA = r"\s*\d{1,4}[-/.]\d{1,2}[-/.]\d{2,4}(?:[ T]\d{1,2}:\d{2}.*)?\s*"
_FMT_ISO = ["%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M:%S",
            "%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%dT%H:%M:%S.%f", "%Y-%m-%dT%H:%M", "%Y/%m/%d",
            "%Y/%m/%d %H:%M:%S"]
_FMT_BR = ["%d/%m/%Y", "%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S", "%d/%m/%y", "%d/%m/%y %H:%M",
           "%d-%m-%Y", "%d-%m-%Y %H:%M:%S", "%d.%m.%Y", "%d.%m.%Y %H:%M:%S"]
_FMT_US = ["%m/%d/%Y", "%m/%d/%Y %H:%M", "%m/%d/%Y %H:%M:%S", "%m/%d/%y", "%m-%d-%Y"]
_FAMILIAS_DATA = {"dd/mm/aaaa": _FMT_ISO + _FMT_BR, "mm/dd/aaaa": _FMT_ISO + _FMT_US}


def _parse_formatos(textos, formatos):
    """Tenta cada formato nos valores (distintos) que ainda não viraram data."""
    def parse_unicos(u):
        valores = _so_textos(u).str.strip().to_numpy(dtype=object)
        res = np.full(len(valores), np.datetime64("NaT", "ns"), dtype="datetime64[ns]")
        for fmt in formatos:
            falta = np.isnat(res) & np.array([isinstance(v, str) for v in valores], dtype=bool)
            if not falta.any():
                break
            conv = pd.to_datetime(pd.Series(valores[falta]), format=fmt, errors="coerce")
            conv = conv.where((conv.dt.year > 1677) & (conv.dt.year < 2262))
            res[falta] = conv.to_numpy(dtype="datetime64[ns]")
        return pd.Series(res)

    if len(textos) > 5000:                     # base grande: só os formatos que deram certo na amostra
        amostra = _amostra_texto(textos, 500)
        taxas = [(_parse_formatos(amostra, [f]).notna().mean(), i, f) for i, f in enumerate(formatos)]
        formatos = [f for taxa, _, f in sorted(taxas, key=lambda t: (-t[0], t[1])) if taxa > 0] or formatos[:1]
    out = _por_unicos(textos.astype(object), parse_unicos)
    return pd.Series(out.to_numpy(dtype="datetime64[ns]"), index=textos.index, name=textos.name)


def _familia_data(textos, pesos=None, minimo=0.9):
    """(família, taxa de acerto ponderada pelas linhas). Datas ambíguas (01/02)
    ficam no padrão BR. minimo: fração que precisa ter cara de data."""
    if textos.empty:
        return None, 0.0
    w = np.ones(len(textos)) if pesos is None else np.asarray(pesos, dtype=float)
    parece = textos.str.fullmatch(_RE_PARECE_DATA).to_numpy(dtype=bool)
    if w[parece].sum() / w.sum() < minimo or not parece.any():
        return None, 0.0
    melhor, taxa_melhor = None, 0.0
    for familia, formatos in _FAMILIAS_DATA.items():
        ok = _parse_formatos(textos, formatos).notna().to_numpy(dtype=bool)
        taxa = w[ok].sum() / w.sum()
        if taxa > taxa_melhor + 1e-9:
            melhor, taxa_melhor = familia, taxa
    return melhor, taxa_melhor


def _sem_fuso(serie):
    if getattr(serie.dt, "tz", None) is not None:
        return serie.dt.tz_localize(None)
    return serie


def converter_data(serie, formato=None):
    """Texto -> data. Aceita ISO (2026-01-31), BR (31/01/2026, 31-01-2026,
    31.01.2026), com ou sem hora. Datas ambíguas (01/02/2026) são lidas como
    dia/mês. O que não for data vira NaT (vazio)."""
    if pd.api.types.is_datetime64_any_dtype(serie):
        return _sem_fuso(serie)
    if formato:
        return _parse_formatos(serie, [formato])
    textos, pesos, _ = _amostra_pesada(serie, 2000)
    familia, _ = _familia_data(textos, pesos, minimo=0.0)
    if familia:
        return _parse_formatos(serie, _FAMILIAS_DATA[familia])
    textos = _so_textos(serie).str.strip()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        try:
            conv = pd.to_datetime(textos, errors="coerce", format="ISO8601", utc=True)
        except (TypeError, ValueError):
            conv = pd.to_datetime(textos, errors="coerce", utc=True, dayfirst=True)
    return conv.dt.tz_localize(None)


def _converter_datas(df, msgs):
    for col in df.columns:
        s = df[col]
        if pd.api.types.is_datetime64_any_dtype(s):
            if getattr(s.dt, "tz", None) is not None:
                df[col] = _sem_fuso(s)
            continue
        if not _eh_texto(s):
            continue
        objetos = s.dropna()
        if len(objetos) and objetos.head(500).map(lambda v: isinstance(v, (datetime, date))).mean() >= 0.9:
            df[col] = pd.to_datetime(s.where(s.map(lambda v: isinstance(v, (datetime, date)))),
                                     errors="coerce")
            continue
        amostra, pesos, _ = _amostra_pesada(s, 2000)
        familia, taxa = _familia_data(amostra, pesos)
        if not familia or taxa < 0.9:
            continue
        conv = _parse_formatos(s, _FAMILIAS_DATA[familia])
        falhou, _ = _falhas(s, conv)
        df[col] = conv
        iso = _parse_formatos(amostra, _FMT_ISO).notna().mean() >= 0.9
        msgs.append(("ok", f"{col}: texto convertido para data ({'aaaa-mm-dd' if iso else familia})."))
        if len(falhou):
            ex = ", ".join(repr(cortar(v, 16)) for v in falhou.drop_duplicates().head(3))
            msgs.append(("aviso", f"{col}: {plural(len(falhou), 'valor não é data válida e virou', 'valores não são datas válidas e viraram')} "
                                  f"vazio (ex: {ex})."))


def preparar(df, verbose=True, preferir_br=True, _info=None):
    """Arruma um DataFrame recém-lido (de arquivo, SQL, clipboard...):
    - tira espaços dos nomes de coluna e células só com espaços;
    - remove linhas e colunas 'Unnamed' totalmente vazias;
    - converte números em texto ('R$ 1.234,56', '12,5%', '(500,00)');
    - converte datas em texto ('31/01/2026', '2026-01-31 14:00').
    Devolve uma CÓPIA. O índice original é mantido (aponta para a linha do arquivo)."""
    info = _info if _info is not None else {"mensagens": []}
    msgs = info.setdefault("mensagens", [])
    df = df.copy()

    nomes, vistos = [], {}
    for c in df.columns:
        nome = c.strip() if isinstance(c, str) else c
        if nome in vistos:
            vistos[nome] += 1
            nome = f"{nome}_{vistos[nome]}"
        else:
            vistos[nome] = 1
        nomes.append(nome)
    if nomes != list(df.columns):
        msgs.append(("info", "Nomes de coluna com espaços nas pontas (ou repetidos) foram ajustados."))
        df.columns = nomes

    for col in df.columns:
        s = df[col]
        if not pd.api.types.is_object_dtype(s) and not pd.api.types.is_string_dtype(s):
            continue
        if pd.api.types.is_object_dtype(s) and pd.api.types.infer_dtype(s, skipna=True) == "mixed":
            if s.dropna().map(lambda v: isinstance(v, (list, dict))).any():
                df[col] = s.map(lambda v: json.dumps(v, ensure_ascii=False, default=str)
                                if isinstance(v, (list, dict)) else v)
                continue
        brancos = [v for v in pd.unique(s.dropna().to_numpy(dtype=object))
                   if isinstance(v, str) and not v.strip()]
        if brancos:
            df[col] = s.where(~s.isin(brancos))

    lixo = [c for c in df.columns if str(c).startswith("Unnamed") and df[c].isna().all()]
    if lixo:
        df = df.drop(columns=lixo)
        msgs.append(("info", f"{len(lixo)} colunas sem nome e vazias foram removidas."))
    vazias = df.isna().all(axis=1)
    if vazias.any():
        df = df[~vazias]
        msgs.append(("info", f"{br(int(vazias.sum()))} linhas totalmente vazias foram removidas."))

    _converter_numeros(df, msgs, preferir_br)
    _converter_datas(df, msgs)
    if verbose:
        for nivel, texto in msgs:
            _status(nivel, texto)
    return df


def _imprimir_carga(info, df):
    partes = [info.get("formato", "DataFrame")]
    if info.get("sep"):
        partes.append(f"separador {info['sep']!r}")
    if info.get("encoding"):
        partes.append(f"encoding {info['encoding']}")
    if info.get("aba"):
        abas = info.get("abas", [])
        extra = f" (outras: {', '.join(a for a in abas if a != info['aba'])}; use --aba)" \
            if len(abas) > 1 else ""
        partes.append(f"aba '{info['aba']}'{extra}")
    if info.get("linha_cabecalho", 1) > 1:
        partes.append(f"cabeçalho na linha {info['linha_cabecalho']}")
    nota(" · ".join(partes))
    nota(f"{br(len(df))} linhas × {len(df.columns)} colunas")
    for nivel, texto in info.get("mensagens", []):
        _status(nivel, texto)


def _arquivos_de_dados(pasta):
    try:
        return sorted(p.name for p in Path(pasta).iterdir()
                      if p.is_file() and p.suffix.lower() in _EXTENSOES)[:12]
    except OSError:
        return []


def _sugerir_arquivos(pasta):
    achados = _arquivos_de_dados(pasta)
    return f"\nArquivos de dados em {pasta}: {', '.join(achados)}" if achados else ""


def _msg_arquivo_inexistente(caminho):
    msg = f"Arquivo não encontrado: {caminho}\nPasta atual: {Path.cwd()}"
    pasta = caminho.parent if caminho.parent.exists() else Path.cwd()
    parecidos = difflib.get_close_matches(caminho.name, _arquivos_de_dados(pasta), n=3, cutoff=0.5)
    if parecidos:
        msg += f"\nVocê quis dizer: {', '.join(parecidos)}?"
    else:
        msg += _sugerir_arquivos(pasta)
    return msg


# ===========================================================================
# 5. QUALIDADE — nulos, duplicatas, regras de negócio, dados pessoais
# ===========================================================================
def diagnostico_nulos(df):
    """Nulos por coluna (quantidade e %). Só mostra quem tem nulo.

    Lembrete: antes de apagar/preencher, pergunte POR QUE está vazio.
    O vazio pode ser informação (ex: pedido ainda não entregue)."""
    qtd = df.isnull().sum()
    out = pd.DataFrame({"nulos": qtd, "pct": (qtd / max(len(df), 1) * 100).round(2)})
    return out[out["nulos"] > 0].sort_values("pct", ascending=False)


def duplicatas(df, subset=None):
    """TODAS as linhas duplicadas (inclusive a primeira ocorrência).
    subset = colunas que definem 'duplicado' (None = linha inteira).
    Para remover: df.drop_duplicates(subset=subset)"""
    out = df[df.duplicated(subset=subset, keep=False)]
    chaves = subset or list(df.columns)
    try:
        return out.sort_values(chaves)
    except TypeError:
        return out


def tratar_nulos(df, metodo="mediana", colunas=None, min_nao_nulos=None):
    """Devolve uma CÓPIA tratada (não altera o original).

    metodo:
      'remover_linhas'   -> dropna(): remove linhas com qualquer nulo
                            (só nas `colunas`, se informadas)
      'remover_colunas'  -> mantém só colunas com pelo menos `min_nao_nulos`
                            valores preenchidos (None = só colunas sem nulo)
      'zero' | 'media' | 'mediana' | 'moda' -> preenche as `colunas`
                            (padrão: numéricas; 'moda' serve também para texto)
    Lembre: fillna(0) não é neutro — o zero entra na média como valor real."""
    metodos = {"remover_linhas", "remover_colunas", "zero", "media", "mediana", "moda"}
    if metodo not in metodos:
        raise ValueError(f"Método desconhecido: {metodo!r}. Use: {', '.join(sorted(metodos))}")
    df = df.copy()
    if metodo == "remover_linhas":
        return df.dropna(subset=colunas)
    if metodo == "remover_colunas":
        if min_nao_nulos is None:
            return df.dropna(axis=1, how="any")
        return df.dropna(axis=1, thresh=min_nao_nulos)
    if colunas is None:
        colunas = [c for c in df.columns if _eh_numerica(df[c])]
    for col in colunas:
        if metodo == "zero":
            df[col] = df[col].fillna(0)
        elif metodo == "media":
            df[col] = df[col].fillna(df[col].mean())
        elif metodo == "mediana":
            df[col] = df[col].fillna(df[col].median())
        else:
            moda = df[col].mode()
            if not moda.empty:
                df[col] = df[col].fillna(moda.iloc[0])
    return df


def padronizar_texto(serie, caixa="upper", acentos=True):
    """Limpa texto para comparar/agrupar: tira espaços das pontas e duplos,
    ajusta a caixa ('upper', 'lower', 'title' ou None) e, com acentos=False,
    remove acentos. ' são  paulo' -> 'SÃO PAULO'."""
    s = _so_textos(serie).str.strip().str.replace(r"\s+", " ", regex=True)
    if not acentos:
        s = s.map(_sem_acento, na_action="ignore")
    if caixa in ("upper", "lower", "title"):
        s = getattr(s.str, caixa)()
    return s.where(s.notna(), serie)


_UFS = {"AC": "acre", "AL": "alagoas", "AP": "amapa", "AM": "amazonas", "BA": "bahia",
        "CE": "ceara", "DF": "distrito federal", "ES": "espirito santo", "GO": "goias",
        "MA": "maranhao", "MT": "mato grosso", "MS": "mato grosso do sul", "MG": "minas gerais",
        "PA": "para", "PB": "paraiba", "PR": "parana", "PE": "pernambuco", "PI": "piaui",
        "RJ": "rio de janeiro", "RN": "rio grande do norte", "RS": "rio grande do sul",
        "RO": "rondonia", "RR": "roraima", "SC": "santa catarina", "SP": "sao paulo",
        "SE": "sergipe", "TO": "tocantins"}


def padronizar_uf(serie):
    """'sp', 'S.P.', 'São Paulo', ' Sao paulo ' -> 'SP'. O que não reconhecer fica como está."""
    mapa = {_normalizar(sigla): sigla for sigla in _UFS}
    mapa.update({_normalizar(nome): sigla for sigla, nome in _UFS.items()})
    return serie.map(lambda v: mapa.get(_normalizar(v), v) if isinstance(v, str) else v)


def _eh_coluna_uf(nome):
    return bool(set(_tokens(nome)) & {"uf", "estado"})


def grafias_inconsistentes(serie, chave=None):
    """Valores que viram o mesmo depois de padronizar (caixa, acento, espaço,
    pontuação): 'SP', 'sp ', 'S.P.' -> um grupo. Em colunas de UF, 'São Paulo'
    também entra no grupo do 'SP'. chave: função própria de padronização.
    Devolve DataFrame com as variantes de cada grupo."""
    contagem = _so_textos(serie).dropna().value_counts()
    if contagem.empty:
        return pd.DataFrame(columns=["padrao", "variantes", "n_variantes", "qtd"])
    if chave is None and _eh_coluna_uf(serie.name):
        chave = lambda v: _normalizar(padronizar_uf(pd.Series([v])).iloc[0])
    chave = chave or _normalizar
    chave = pd.Series([chave(v) for v in contagem.index], index=contagem.index)
    grupos = []
    for k, idx in chave.groupby(chave).groups.items():
        if len(idx) > 1 and k:
            variantes = contagem[list(idx)].sort_values(ascending=False)
            grupos.append({"padrao": variantes.index[0], "variantes": list(variantes.index),
                           "n_variantes": len(variantes), "qtd": int(variantes.sum())})
    out = pd.DataFrame(grupos, columns=["padrao", "variantes", "n_variantes", "qtd"])
    return out.sort_values("qtd", ascending=False).reset_index(drop=True)


# --- dados pessoais (LGPD) ---------------------------------------------------
def _documento_texto(v):
    """Valor de CPF/CNPJ como texto: '529.982.247-25', 52998224725 e 52998224725.0 servem."""
    if isinstance(v, str):
        return v
    if isinstance(v, (int, np.integer)) and not isinstance(v, (bool, np.bool_)):
        return str(int(v))
    if isinstance(v, (float, np.floating)) and math.isfinite(v) and float(v).is_integer():
        return str(int(v))
    return None


def _so_digitos(serie, tamanho):
    """CPF/CNPJ como string só de dígitos/letras, com zeros à esquerda recolocados."""
    s = serie.astype(object).map(_documento_texto).astype("string")
    s = s.str.upper().str.replace(r"[^0-9A-Z]", "", regex=True)
    return s.where(s.str.len() > 0).str.zfill(tamanho)


def _validar_documento(serie, tamanho):
    """Dígitos verificadores de CPF (11) ou CNPJ (14), calculado só nos valores distintos."""
    out = _por_unicos(serie.astype(object), lambda u: _validar_unicos(u, tamanho))
    return _bool(out)


def _validar_unicos(serie, tamanho):
    s = _so_digitos(serie, tamanho)
    valores = s.to_numpy(dtype=object, na_value=None)
    padrao = re.compile(r"\d{11}" if tamanho == 11 else r"[0-9A-Z]{12}\d{2}")
    mascara = np.array([isinstance(v, str) and bool(padrao.fullmatch(v)) for v in valores], dtype=bool)
    out = np.zeros(len(valores), dtype=bool)
    if mascara.any():
        arr = np.frombuffer("".join(valores[mascara]).encode("ascii"), dtype=np.uint8)
        arr = arr.reshape(-1, tamanho).astype(int) - 48          # '0'->0 ... 'A'->17 (CNPJ alfanumérico)
        if tamanho == 11:
            d1 = (arr[:, :9] * np.arange(10, 1, -1)).sum(1) * 10 % 11 % 10
            d2 = (arr[:, :10] * np.arange(11, 1, -1)).sum(1) * 10 % 11 % 10
            ok = (d1 == arr[:, 9]) & (d2 == arr[:, 10])
        else:
            p1 = np.array([5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2])
            p2 = np.array([6, 5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2])
            r1 = (arr[:, :12] * p1).sum(1) % 11
            d1 = np.where(r1 < 2, 0, 11 - r1)
            r2 = (np.column_stack([arr[:, :12], d1]) * p2).sum(1) % 11
            d2 = np.where(r2 < 2, 0, 11 - r2)
            ok = (d1 == arr[:, 12]) & (d2 == arr[:, 13])
        repetido = (arr == arr[:, [0]]).all(axis=1)
        out[mascara] = ok & ~repetido
    return pd.Series(out, index=serie.index)


def validar_cpf(valor):
    """True se o CPF for válido (dígitos verificadores). Aceita série ou valor único,
    com ou sem pontuação. 111.111.111-11 e parecidos são inválidos."""
    if isinstance(valor, pd.Series):
        return _validar_documento(valor, 11)
    return bool(_validar_documento(pd.Series([valor], dtype=object), 11).iloc[0])


def validar_cnpj(valor):
    """True se o CNPJ for válido. Aceita o CNPJ alfanumérico (julho/2026 em diante)."""
    if isinstance(valor, pd.Series):
        return _validar_documento(valor, 14)
    return bool(_validar_documento(pd.Series([valor], dtype=object), 14).iloc[0])


_PII_NOMES = {"cpf": "CPF", "cnpj": "CNPJ", "email": "e-mail", "mail": "e-mail",
              "telefone": "telefone", "celular": "telefone", "fone": "telefone",
              "whatsapp": "telefone", "rg": "RG", "pis": "PIS", "nis": "PIS"}
_RE_EMAIL = r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+"
_RE_CPF_FMT = r"\d{3}\.\d{3}\.\d{3}-\d{2}"
_RE_CNPJ_FMT = r"[0-9A-Za-z]{2}\.[0-9A-Za-z]{3}\.[0-9A-Za-z]{3}/[0-9A-Za-z]{4}-\d{2}"
_RE_TELEFONE = r"(?:\+?55\s?)?\(?\d{2}\)?\s?9?\d{4}[-\s]\d{4}"


def detectar_pii(df):
    """Colunas com dados pessoais (LGPD): {coluna: 'CPF'|'CNPJ'|'e-mail'|'telefone'|'nome'|...}.
    Usa o nome da coluna e o formato dos valores."""
    out = {}
    for col in df.columns:
        s = df[col]
        if pd.api.types.is_datetime64_any_dtype(s) or pd.api.types.is_bool_dtype(s):
            continue
        toks = set(_tokens(col))
        por_nome = next((_PII_NOMES[t] for t in _tokens(col) if t in _PII_NOMES), None)
        if por_nome:
            out[col] = por_nome
            continue
        if not _eh_texto(s):
            continue
        amostra = _amostra_texto(s, 300)
        if amostra.empty:
            continue
        if amostra.str.fullmatch(_RE_EMAIL).mean() >= 0.8:
            out[col] = "e-mail"
        elif amostra.str.fullmatch(_RE_CPF_FMT).mean() >= 0.8:
            out[col] = "CPF"
        elif amostra.str.fullmatch(_RE_CNPJ_FMT).mean() >= 0.8:
            out[col] = "CNPJ"
        elif amostra.str.fullmatch(r"\d{11}").mean() >= 0.9 and validar_cpf(amostra).mean() >= 0.8:
            out[col] = "CPF"
        elif amostra.str.fullmatch(r"\d{14}").mean() >= 0.9 and validar_cnpj(amostra).mean() >= 0.8:
            out[col] = "CNPJ"
        elif amostra.str.fullmatch(_RE_TELEFONE).mean() >= 0.8:
            out[col] = "telefone"
        elif "nome" in toks and not toks & {"produto", "item", "arquivo", "empresa", "loja",
                                            "fantasia", "razao", "social", "categoria"}:
            out[col] = "nome"
    return out


def _mascarar_valor(v, tipo):
    if not isinstance(v, str) and _na(v):
        return v
    v = str(v)
    if tipo == "CPF":
        d = re.sub(r"\D", "", v).zfill(11)
        return f"***.{d[3:6]}.{d[6:9]}-**"
    if tipo == "e-mail" and "@" in v:
        usuario, dominio = v.split("@", 1)
        return usuario[:2] + "***@" + dominio
    if tipo == "telefone":
        d = re.sub(r"\D", "", v)
        return "(**) ****-" + d[-4:] if len(d) >= 4 else "****"
    if tipo == "nome":
        partes = v.split()
        return partes[0] + (f" {partes[-1][0]}." if len(partes) > 1 else "")
    return v[:2] + "*" * max(len(v) - 4, 3) + v[-2:] if len(v) > 4 else "****"


def mascarar_pii(df, colunas=None):
    """CÓPIA com dados pessoais mascarados (para compartilhar ou printar):
    CPF ***.456.789-**, e-mail jo***@x.com, telefone (**) ****-4321, nome 'João S.'.
    colunas: dict {coluna: tipo} ou lista (None = detectar_pii). CNPJ não é mascarado."""
    if colunas is None:
        colunas = detectar_pii(df)
    elif not isinstance(colunas, dict):
        detectadas = detectar_pii(df)
        colunas = {c: detectadas.get(c, "outro") for c in colunas}
    out = df.copy()
    for col, tipo in colunas.items():
        if tipo != "CNPJ":
            out[col] = df[col].map(lambda v, t=tipo: _mascarar_valor(v, t)).astype(object)
    return out


def checar_qualidade(df, col_valor=None, pii=None):
    """Bateria de checagens automáticas de regras de negócio. Devolve DataFrame
    (nivel, checagem, coluna, detalhe, qtd) só com o que encontrou de estranho."""
    pii = detectar_pii(df) if pii is None else pii
    achados = []

    def add(nivel, checagem, coluna, detalhe, qtd=None):
        achados.append({"nivel": nivel, "checagem": checagem, "coluna": coluna,
                        "detalhe": detalhe, "qtd": qtd})

    hoje = pd.Timestamp.today().normalize() + pd.Timedelta(days=1)
    for col in df.columns:
        s = df[col]
        preenchidos = s.dropna()
        if preenchidos.empty:
            add("aviso", "coluna vazia", col, "nenhum valor preenchido", len(s))
            continue
        n_dist = _nunique(s)
        if n_dist == 1 and len(preenchidos) > 1:
            add("info", "coluna constante", col,
                f"sempre {_celula(preenchidos.iloc[0], max_cel=20)!s}; não ajuda a explicar nada")
        if _parece_codigo_nome(col) and col not in pii and len(preenchidos) >= 10:
            rep = int(preenchidos.duplicated().sum())
            if 0 < rep and n_dist >= 0.95 * len(preenchidos):
                add("aviso", "ID repetido", col,
                    f"{br(rep)} valores repetidos numa coluna quase única; se for a chave da "
                    "tabela, há registros duplicados", rep)
        if pd.api.types.is_datetime64_any_dtype(s):
            if not set(_tokens(col)) & _DATA_NAO_EVENTO:
                futuras = preenchidos[preenchidos > hoje]
                if len(futuras):
                    add("aviso", "data no futuro", col,
                        f"{br(len(futuras))} datas depois de hoje (a maior: {data_br(futuras.max())})",
                        len(futuras))
            antigas = preenchidos[preenchidos < pd.Timestamp("1900-01-01")]
            if len(antigas):
                add("aviso", "data implausível", col,
                    f"{br(len(antigas))} datas antes de 1900 (a menor: {data_br(antigas.min())})",
                    len(antigas))
        if _eh_texto(s):
            contagem = _so_textos(s).value_counts()
            textos = pd.Series(contagem.index.to_numpy(dtype=object))
            if 2 <= n_dist <= 5000:
                g = grafias_inconsistentes(s)
                if not g.empty:
                    exemplos = "; ".join(" / ".join(repr(v) for v in variantes[:3])
                                         for variantes in g["variantes"].head(2))
                    funcao = "padronizar_uf()" if _eh_coluna_uf(col) else "padronizar_texto()"
                    add("aviso", "grafias diferentes", col,
                        f"{plural(len(g), 'valor escrito', 'valores escritos')} de mais de um jeito "
                        f"({exemplos}); agrupamentos ficam divididos — use {funcao}",
                        int(g["qtd"].sum()))
            sobra = (textos.str.strip().ne(textos) | textos.str.contains("  ", regex=False)).to_numpy(bool)
            n_sobra = int(contagem.to_numpy()[sobra].sum())
            if n_sobra:
                add("info", "espaços sobrando", col,
                    f"{br(n_sobra)} valores com espaço nas pontas ou duplo (atrapalha "
                    "filtro e merge)", n_sobra)
            if col not in pii:
                amostra = _amostra_texto(s, 2000)
                if len(amostra) >= 5:
                    parece = amostra.str.fullmatch(_PADRAO_BR) | amostra.str.fullmatch(_PADRAO_US)
                    if 0.5 <= parece.mean() < 0.95:
                        intrusos = amostra[~parece].head(3)
                        add("aviso", "número guardado como texto", col,
                            f"{pct(parece.mean() * 100, 0)} parecem número, mas há textos "
                            f"misturados (ex: {', '.join(repr(cortar(v, 12)) for v in intrusos)}); "
                            "use converter_numero()")
        if pii.get(col) in ("CPF", "CNPJ"):
            validos = (validar_cpf if pii[col] == "CPF" else validar_cnpj)(preenchidos)
            invalidos = preenchidos[~validos.to_numpy()]
            if len(invalidos):
                ex = ", ".join(repr(cortar(str(v), 18)) for v in invalidos.drop_duplicates().head(3))
                add("alerta", f"{pii[col]} inválido", col,
                    f"{br(len(invalidos))} valores ({pct(len(invalidos) / len(preenchidos) * 100)}) "
                    f"com dígito verificador errado (ex: {ex})", len(invalidos))
    if col_valor is not None and col_valor in df.columns and _eh_numerica(df[col_valor]):
        v = _num(df[col_valor])
        neg, zeros = int((v < 0).sum()), int((v == 0).sum())
        if neg:
            add("info", "valores negativos", col_valor,
                f"{br(neg)} valores negativos somando {br(v[v < 0].sum())} (estornos? débitos?)", neg)
        if zeros:
            add("info", "valores zerados", col_valor, f"{br(zeros)} valores iguais a zero", zeros)
    if pii:
        tipos = ", ".join(f"{c} ({t})" for c, t in pii.items())
        add("info", "dados pessoais", ", ".join(map(str, pii)),
            f"LGPD: {tipos}. Para compartilhar, use mascarar_pii(df).")
    return pd.DataFrame(achados, columns=["nivel", "checagem", "coluna", "detalhe", "qtd"])


# ===========================================================================
# 6. ESTATÍSTICA
# ===========================================================================
def resumo_numerico(serie):
    """Todos os indicadores de uma coluna numérica de uma vez (pd.Series)."""
    s = _num(serie).dropna()
    contagem = s.value_counts()
    repete = not contagem.empty and contagem.iloc[0] > 1
    media, desvio = s.mean(), s.std()
    q1, q3 = (s.quantile(0.25), s.quantile(0.75)) if len(s) else (np.nan, np.nan)
    return pd.Series(dtype=object, data={
        "contagem": int(s.count()),
        "nulos": int(serie.isna().sum()),
        "soma": s.sum(),
        "media": media,
        "mediana": s.median(),
        "moda": contagem.index[0] if repete else np.nan,
        "moda_freq": int(contagem.iloc[0]) if repete else 0,
        "desvio_padrao": desvio,
        "variancia": s.var(),
        "coef_variacao_pct": desvio / abs(media) * 100 if media else np.nan,
        "min": s.min(),
        "q25": q1,
        "q75": q3,
        "max": s.max(),
        "amplitude": s.max() - s.min(),
        "iqr": q3 - q1,
        "negativos": int((s < 0).sum()),
        "zeros": int((s == 0).sum()),
    })


def limites_iqr(serie, k=1.5):
    """(limite inferior, limite superior) = Q1 - k*IQR, Q3 + k*IQR.
    k=1.5 é a regra clássica; k=3 pega só os extremos."""
    s = _num(serie)
    q1, q3 = s.quantile([0.25, 0.75])
    iqr = q3 - q1
    return q1 - k * iqr, q3 + k * iqr


def outliers_iqr(df, coluna, k=1.5):
    """Linhas com valor fora dos limites do IQR. Candidatos a erro ou fraude."""
    inf, sup = limites_iqr(df[coluna], k)
    v = _num(df[coluna])
    return df[(v < inf) | (v > sup)]


def _assimetrica_positiva(serie):
    """Valores quase todos positivos e com cauda longa à direita (típico de
    dinheiro): nesses casos, o 'atípico' deve ser medido na escala log."""
    s = _num(serie).dropna()
    pos = s[s > 0]
    return len(pos) >= 0.9 * max(len(s), 1) and len(pos) > 0 and pos.mean() > 1.5 * pos.median()


def zscore_robusto(serie, log=False):
    """Z-score modificado (Iglewicz & Hoaglin): 0,6745·(x − mediana) / MAD.
    |z| > 3,5 = atípico. Diferente do z comum, não é distorcido pelos próprios
    outliers. Se MAD = 0, usa o desvio absoluto médio.
    log=True mede na escala log10(|x|) — melhor para valores em dinheiro, que
    têm cauda longa (senão metade da cauda vira 'atípica'); zeros ficam NaN."""
    s = _num(serie)
    if log:
        a = s.abs()
        s = np.log10(a.where(a > 0))
    med = s.median()
    mad = (s - med).abs().median()
    if mad and not _na(mad):
        return 0.6745 * (s - med) / mad
    desvio_medio = (s - med).abs().mean()
    if desvio_medio and not _na(desvio_medio):
        return (s - med) / (1.253314 * desvio_medio)
    return (s - med) * 0.0


def outliers_mad(df, coluna, limite=3.5, log=None):
    """Linhas com |z-score robusto| > limite, com a coluna 'z_robusto', da mais
    atípica para a menos. log=None decide sozinho (log se o valor for assimétrico)."""
    log = _assimetrica_positiva(df[coluna]) if log is None else log
    z = zscore_robusto(df[coluna], log=log)
    out = df[z.abs() > limite].copy()
    out["z_robusto"] = z[z.abs() > limite]
    return out.reindex(out["z_robusto"].abs().sort_values(ascending=False).index)


def top_n(df, coluna, n=5):
    """As N maiores linhas por coluna (para as menores: df.nsmallest(n, coluna))."""
    return df.nlargest(n, coluna)


def correlacao(df, colunas=None, metodo="pearson"):
    """Matriz de correlação entre colunas numéricas.
    metodo: 'pearson' (linear) ou 'spearman' (ordem; melhor com outliers)."""
    base = df[colunas] if colunas else df.select_dtypes("number")
    return base.corr(method=metodo)


def pares_correlacao(df, colunas=None, metodo="pearson", minimo=0.3):
    """Pares de colunas com |correlação| >= minimo, do mais forte para o mais fraco."""
    m = correlacao(df, colunas, metodo)
    linhas = []
    cols = list(m.columns)
    for i, a in enumerate(cols):
        for b in cols[i + 1:]:
            r = m.loc[a, b]
            if not _na(r) and abs(r) >= minimo:
                forca = "forte" if abs(r) >= 0.7 else "moderada" if abs(r) >= 0.4 else "fraca"
                linhas.append({"coluna_a": a, "coluna_b": b, "r": r, "forca": forca,
                               "direcao": "positiva" if r > 0 else "negativa"})
    out = pd.DataFrame(linhas, columns=["coluna_a", "coluna_b", "r", "forca", "direcao"])
    return out.reindex(out["r"].abs().sort_values(ascending=False).index).reset_index(drop=True)


def _passo_bonito(bruto):
    if bruto <= 0 or _na(bruto):
        return 1.0
    exp = 10 ** math.floor(math.log10(bruto))
    for m in (1, 2, 2.5, 5, 10):
        if bruto <= m * exp:
            return m * exp
    return 10 * exp


def histograma(serie, faixas=10, cortar_outliers=True):
    """Distribuição em faixas de largura 'redonda'. Com cortar_outliers, os
    valores fora do IQR viram duas linhas à parte (para não achatar o gráfico).
    Devolve DataFrame: faixa, de, ate, qtd, pct, fora (bool)."""
    s = _num(serie).dropna()
    colunas = ["faixa", "de", "ate", "qtd", "pct", "fora"]
    if s.empty:
        return pd.DataFrame(columns=colunas)
    lo, hi = s.min(), s.max()
    if cortar_outliers:
        inf, sup = limites_iqr(s)
        if sup > inf:
            lo, hi = max(lo, inf), min(hi, sup)
    linhas = []
    if bool((s % 1 == 0).all()) and hi - lo <= 20:
        nucleo = s[(s >= lo) & (s <= hi)]
        lo, hi, casas = int(nucleo.min()), int(nucleo.max()), 0
        for v in range(int(lo), int(hi) + 1):
            linhas.append([br(v), v, v, int((s == v).sum()), False])
    else:
        passo = _passo_bonito((hi - lo) / faixas) if hi > lo else 1.0
        bordas = [math.floor(lo / passo) * passo]
        while bordas[-1] < hi or len(bordas) < 2:
            bordas.append(bordas[-1] + passo)
        lo, hi = bordas[0], bordas[-1]                 # a faixa mostrada manda: nada fica de fora
        contagem, _ = np.histogram(s[(s >= lo) & (s <= hi)], bins=bordas)
        casas = 0 if passo >= 1 and float(passo).is_integer() else _casas(pd.Series([passo]))
        for a, b, q in zip(bordas[:-1], bordas[1:], contagem):
            linhas.append([f"{br(a, casas)} a {br(b, casas)}", a, b, int(q), False])
    abaixo, acima = s[s < lo], s[s > hi]
    if len(abaixo):
        linhas.insert(0, [f"abaixo de {br(float(lo), casas)}", abaixo.min(), lo, len(abaixo), True])
    if len(acima):
        linhas.append([f"acima de {br(float(hi), casas)}", hi, acima.max(), len(acima), True])
    out = pd.DataFrame(linhas, columns=["faixa", "de", "ate", "qtd", "fora"])
    out["pct"] = out["qtd"] / len(s) * 100
    return out[colunas]


def leitura_em_frase(serie, n_outliers=None, casas=None):
    """Traduz os números em texto — a frase é o produto, não a tabela."""
    s = _num(serie).dropna()
    if s.empty:
        return "Não há valores numéricos para ler."
    c = _casas(s) if casas is None else casas
    r = resumo_numerico(s)
    media, med, dp, cv = r["media"], r["mediana"], r["desvio_padrao"], r["coef_variacao_pct"]
    simetrica = bool(med) and abs(media - med) <= 0.25 * abs(med) and not r["negativos"]
    if len(s) > 1 and not _na(dp) and simetrica:
        frases = [f"O valor típico (mediana) é {br(med, c)}, com oscilação de cerca de "
                  f"{br(dp, c)} para mais ou para menos."]
    else:
        frases = [f"O valor típico (mediana) é {br(med, c)}."]
    if med:
        dif = (media - med) / abs(med)
        if dif > 0.10:
            frases.append(f"A média ({br(media, c)}) fica {pct(dif * 100, 0)} acima da mediana: "
                          "poucos valores altos puxam a média para cima — para falar do "
                          "'típico', prefira a mediana.")
        elif dif < -0.10:
            frases.append(f"A média ({br(media, c)}) fica {pct(-dif * 100, 0)} abaixo da mediana: "
                          "poucos valores baixos puxam a média para baixo.")
        else:
            frases.append("Média e mediana são próximas: a distribuição é equilibrada.")
    if r["negativos"]:
        frases.append(f"Há {br(r['negativos'])} valores negativos "
                      f"({pct(r['negativos'] / len(s) * 100)}) — estornos, devoluções ou débitos?")
    elif not _na(cv):
        nivel = ("baixa" if cv < 15 else "moderada" if cv < 30 else "alta" if cv < 100
                 else "muito alta")
        frases.append(f"A variação é {nivel} (coeficiente de variação de {pct(cv)}).")
    frases.append(f"Metade dos valores fica entre {br(r['q25'], c)} e {br(r['q75'], c)}.")
    if r["negativos"] == 0 and r["soma"] > 0 and len(s) >= 20:
        n10 = max(1, int(round(len(s) * 0.10)))
        conc = s.nlargest(n10).sum() / r["soma"]
        if conc >= 0.30:
            frases.append(f"Os 10% maiores valores concentram {pct(conc * 100, 0)} da soma.")
    if n_outliers is None:
        inf, sup = limites_iqr(s)
        n_outliers = int(((s < inf) | (s > sup)).sum())
    if n_outliers:
        frases.append(f"{br(n_outliers)} valores ({pct(n_outliers / len(s) * 100)}) fogem do "
                      "padrão (outliers) e merecem conferência.")
    return " ".join(frases)


# ===========================================================================
# 7. AGRUPAMENTO, CRUZAMENTO E TEMPO
# ===========================================================================
def agrupar(df, chave, valor, funcs=("sum", "mean", "count")):
    """groupby com várias métricas, ordenado pela primeira. Vazios viram um grupo.
    chave pode ser lista: agrupar(df, ['Produto', 'Vendedor'], 'Total')"""
    funcs = list(funcs)
    g = df.groupby(chave, dropna=False, observed=True)[valor].agg(funcs)
    return g.sort_values(funcs[0], ascending=False)


def cruzar(df, linhas, colunas, valor, func="sum"):
    """Tabela dinâmica (estilo Excel) com totais nas margens."""
    return df.pivot_table(index=linhas, columns=colunas, values=valor, aggfunc=func,
                          margins=True, margins_name="Total", fill_value=0, observed=True)


def curva_abc(df, chave, valor=None, a=80, b=95):
    """Curva ABC / Pareto: cada grupo com qtd, total, % do total, % acumulado e
    classe (A = os poucos que fazem ~80%, B = até 95%, C = o resto).
    Sem `valor`, usa a contagem de linhas."""
    k = df[chave].astype(object).where(df[chave].notna(), "(vazio)")
    if valor is None:
        g = k.value_counts().to_frame("qtd")
        base = "qtd"
    else:
        tmp = pd.DataFrame({"k": k, "v": _num(df[valor])})
        g = tmp.groupby("k", sort=False)["v"].agg(qtd="size", total="sum", media="mean",
                                                  mediana="median")
        base = "total"
    g = g.sort_values(base, ascending=False)
    total = g[base].sum()
    g["pct"] = g[base] / total * 100 if total else np.nan
    g["pct_acum"] = g["pct"].cumsum()
    antes = g["pct_acum"] - g["pct"]
    g["classe"] = np.where(antes < a, "A", np.where(antes < b, "B", "C"))
    g.index.name = chave
    return g


def resumo_por_grupo(df, chave, valor=None, limite=15):
    """curva_abc() resumida: os `limite` maiores grupos + uma linha '(outros N)'."""
    g = curva_abc(df, chave, valor)
    if len(g) <= limite:
        return g
    topo, resto = g.head(limite), g.iloc[limite:]
    outros = {"qtd": resto["qtd"].sum(), "pct": resto["pct"].sum(), "pct_acum": 100.0,
              "classe": ""}
    if valor is not None:
        outros.update(total=resto["total"].sum(), media=resto["total"].sum() / resto["qtd"].sum(),
                      mediana=np.nan)
    linha = pd.DataFrame([outros], index=[f"(outros {len(resto)})"])
    return pd.concat([topo, linha[topo.columns]])


def frequencia(df, coluna, limite=None):
    """Contagem de cada valor (vazios incluídos), com % e % acumulado."""
    vc = df[coluna].astype(object).where(df[coluna].notna(), "(vazio)").value_counts()
    out = vc.to_frame("qtd")
    out["pct"] = out["qtd"] / out["qtd"].sum() * 100
    out["pct_acum"] = out["pct"].cumsum()
    out.index.name = coluna
    return out.head(limite) if limite else out


_FREQS = {"D": ("D", {}), "W": ("W-MON", {"label": "left", "closed": "left"}),
          "MS": ("MS", {}), "QS": ("QS", {}), "YS": ("YS", {})}


def _freq_auto(dias):
    return "QS" if dias > 3 * 365 else "MS" if dias > 92 else "W" if dias > 21 else "D"


def _rotulo_periodo(ts, freq):
    if freq == "MS":
        return f"{_MESES[ts.month - 1]}/{ts.year}"
    if freq == "QS":
        return f"{(ts.month - 1) // 3 + 1}º tri/{ts.year}"
    if freq == "YS":
        return str(ts.year)
    if freq == "W":
        return f"sem. {ts.strftime('%d/%m/%y')}"
    return f"{ts.strftime('%d/%m/%Y')} {_DIAS[ts.dayofweek]}"


def serie_temporal(df, col_data, valor=None, freq="MS", func="sum"):
    """Agrega por período. freq: 'D' dia, 'W' semana (começa na segunda),
    'MS' mês, 'QS' trimestre, 'YS' ano. Sem `valor`, conta linhas."""
    d = converter_data(df[col_data])
    regra, kw = _FREQS.get(freq, (freq, {}))
    v = pd.Series(1.0, index=df.index) if valor is None else _num(df[valor])
    tmp = pd.Series(v.to_numpy(), index=pd.DatetimeIndex(d)).loc[d.notna().to_numpy()]
    return tmp.resample(regra, **kw).agg("sum" if valor is None else func)


def resumo_temporal(df, col_data, valor=None, freq=None):
    """Tabela por período: qtd, total, variação % e se o período está incompleto
    (dados que começam/terminam no meio do mês não dão variação justa)."""
    d = converter_data(df[col_data])
    validas = d.dropna()
    colunas = ["periodo", "rotulo", "qtd", "total", "var_pct", "parcial"]
    if validas.empty:
        return pd.DataFrame(columns=colunas), freq
    ini, fim = validas.min(), validas.max()
    freq = freq or _freq_auto((fim - ini).days)
    regra, kw = _FREQS.get(freq, (freq, {}))
    v = np.ones(len(df)) if valor is None else _num(df[valor]).to_numpy()
    base = pd.DataFrame({"qtd": 1, "total": v}, index=pd.DatetimeIndex(d))[d.notna().to_numpy()]
    r = base.resample(regra, **kw).agg({"qtd": "sum", "total": "sum"})
    r["qtd"] = r["qtd"].astype(int)
    medida = r["qtd"] if valor is None else r["total"]
    var = (medida / medida.shift(1) - 1) * 100
    r["var_pct"] = var.replace([np.inf, -np.inf], np.nan)
    tolerancia = {"MS": 2, "QS": 5, "YS": 10, "W": 0, "D": 0}.get(freq, 0)
    r["parcial"] = False
    if len(r) and freq != "D":
        if ini.normalize() > r.index[0] + pd.Timedelta(days=tolerancia):
            r.iloc[0, r.columns.get_loc("parcial")] = True
        fim_ultimo = r.index[-1] + pd.tseries.frequencies.to_offset(regra) - pd.Timedelta(days=1)
        if fim.normalize() < fim_ultimo - pd.Timedelta(days=tolerancia):
            r.iloc[-1, r.columns.get_loc("parcial")] = True
    r["rotulo"] = [_rotulo_periodo(ts, freq) for ts in r.index]
    r["periodo"] = r.index
    if valor is None:
        r["total"] = np.nan
    return r[colunas].reset_index(drop=True), freq


def coorte(df, col_entidade, col_data, freq="M", pct=True):
    """Matriz de coorte (retenção): linhas = mês da 1ª aparição da entidade,
    colunas = períodos depois disso, valores = % (ou qtd) de entidades ativas."""
    d = converter_data(df[col_data])
    tmp = pd.DataFrame({"ent": df[col_entidade], "per": d.dt.to_period(freq)}).dropna()
    tmp["ord"] = pd.PeriodIndex(tmp["per"]).asi8
    tmp["inicio"] = tmp.groupby("ent")["ord"].transform("min")
    tmp["idade"] = tmp["ord"] - tmp["inicio"]
    tmp["coorte"] = tmp.groupby("ent")["per"].transform("min").astype(str)
    m = tmp.groupby(["coorte", "idade"])["ent"].nunique().unstack(fill_value=0)
    if pct:
        m = m.div(m[0], axis=0) * 100
    m.columns.name = "períodos depois"
    return m


def _serie_horas(df, col_data):
    """Hora de cada registro: da própria coluna de data (se tiver hora) ou de uma
    coluna separada tipo 'Hora' ('14:35'). None se não houver hora."""
    d = df[col_data]
    if pd.api.types.is_datetime64_any_dtype(d):
        validas = d.dropna()
        if len(validas) and ((validas.dt.hour != 0) | (validas.dt.minute != 0)).any():
            return d.dt.hour.astype("float64")
    for col in df.columns:
        if col == col_data or not set(_tokens(col)) & {"hora", "horario", "hr", "time"}:
            continue
        textos = _so_textos(df[col]).str.strip()
        hh = textos.str.extract(r"^(\d{1,2}):\d{2}", expand=False)
        if hh.notna().sum() >= 0.9 * max(textos.notna().sum(), 1) and hh.notna().any():
            return pd.to_numeric(hh, errors="coerce")
    return None


# ===========================================================================
# 8. ANTIFRAUDE — Benford, atípicos, duplicidade, fracionamento, calendário
# ===========================================================================
def _chi2_sf(x, gl):
    """P(Qui² > x). Exato para gl par; aproximação de Wilson-Hilferty para ímpar."""
    if x <= 0:
        return 1.0
    if gl % 2 == 0:
        termo, soma = 1.0, 1.0
        for i in range(1, gl // 2):
            termo *= (x / 2) / i
            soma += termo
        return math.exp(-x / 2) * soma
    z = ((x / gl) ** (1 / 3) - (1 - 2 / (9 * gl))) / math.sqrt(2 / (9 * gl))
    return 0.5 * math.erfc(z / math.sqrt(2))


def benford(serie, digitos=1):
    """Teste da Lei de Benford no 1º dígito (digitos=1) ou nos 2 primeiros (2).
    Devolve DataFrame (digito, qtd, observado, esperado, diferenca, z) com o
    resumo em .attrs: n, mad, conformidade, chi2, p_valor, ordens.
    Conformidade pelo MAD (Nigrini): próxima, aceitável, marginal, não conforme."""
    v = _num(serie).dropna().abs()
    v = v[v >= 10] if digitos == 2 else v[v > 0]
    faixa = np.arange(1, 10) if digitos == 1 else np.arange(10, 100)
    esperado = np.log10(1 + 1 / faixa)
    if v.empty:
        out = pd.DataFrame({"digito": faixa, "qtd": 0, "observado": 0.0, "esperado": esperado,
                            "diferenca": -esperado, "z": 0.0})
        out.attrs.update(n=0, mad=np.nan, conformidade="sem dados", chi2=np.nan, p_valor=np.nan,
                         ordens=0.0)
        return out
    expoente = np.floor(np.log10(v.to_numpy())) - (digitos - 1)
    d = np.floor(v.to_numpy() / 10.0 ** expoente + 1e-9).astype(int)
    contagem = pd.Series(d).value_counts().reindex(faixa, fill_value=0).to_numpy()
    n = int(contagem.sum())
    obs = contagem / n
    dif = obs - esperado
    ep = np.sqrt(esperado * (1 - esperado) / n)
    corr = 1 / (2 * n)
    z = np.where(np.abs(dif) > corr, (np.abs(dif) - corr), np.abs(dif)) / ep
    mad = float(np.abs(dif).mean())
    cortes = (0.006, 0.012, 0.015) if digitos == 1 else (0.0012, 0.0018, 0.0022)
    conformidade = ("próxima" if mad <= cortes[0] else "aceitável" if mad <= cortes[1]
                    else "marginal" if mad <= cortes[2] else "não conforme")
    chi2 = float(n * (dif ** 2 / esperado).sum())
    q01, q99 = np.quantile(v, [0.01, 0.99])
    out = pd.DataFrame({"digito": faixa, "qtd": contagem, "observado": obs, "esperado": esperado,
                        "diferenca": dif, "z": z})
    out.attrs.update(n=n, mad=mad, conformidade=conformidade, chi2=chi2,
                     p_valor=_chi2_sf(chi2, len(faixa) - 1),
                     ordens=float(np.log10(q99 / q01)) if q01 > 0 else 0.0)
    return out


def valores_repetidos(serie, top=10):
    """Valores exatos que mais se repetem (teste de duplicação de números)."""
    s = _num(serie).dropna()
    vc = s.value_counts()
    vc = vc[vc > 1].head(top)
    out = pd.DataFrame({"valor": vc.index, "vezes": vc.to_numpy()})
    out["pct"] = out["vezes"] / max(len(s), 1) * 100
    out["soma"] = out["valor"] * out["vezes"]
    return out


def valores_redondos(serie):
    """Fatia de valores 'redondos': múltiplos de 1.000 (entre os >= 1.000), de
    100 (entre os >= 100) e sem centavos. Excesso de redondos sugere estimativa
    ou valor inventado — ou só um negócio com preços cheios."""
    s = _num(serie).dropna().abs().round(2)

    def fatia(mult):
        base = s[s >= mult]
        return (base % mult == 0).mean() * 100 if len(base) else np.nan

    return {"mult_1000": fatia(1000), "mult_100": fatia(100),
            "sem_centavos": (s % 1 == 0).mean() * 100 if len(s) else np.nan,
            "n_1000": int((s >= 1000).sum()), "n_100": int((s >= 100).sum())}


def possiveis_duplicatas(df, col_valor, col_data=None, col_entidade=None, outras=None):
    """Possível pagamento em duplicidade: mesmo valor + mesma entidade + mesmo dia
    (sem entidade: mesmo valor no mesmo minuto). Devolve as linhas, com 'grupo_dup'."""
    chaves = pd.DataFrame({"_v": _num(df[col_valor]).round(2)}, index=df.index)
    if col_entidade:
        chaves["_e"] = df[col_entidade].astype(object)
    if col_data:
        d = converter_data(df[col_data])
        chaves["_d"] = d.dt.normalize() if col_entidade else d.dt.floor("min")
    for c in outras or []:
        chaves[c] = df[c].astype(object)
    if chaves.shape[1] == 1:
        raise ValueError("Informe col_data e/ou col_entidade: só o valor igual não basta.")
    completas = chaves.notna().all(axis=1) & (chaves["_v"] != 0)
    dup = chaves.duplicated(keep=False) & completas
    out = df[dup].copy()
    out["grupo_dup"] = chaves[dup].groupby(list(chaves.columns), sort=False).ngroup()
    return out.sort_values("grupo_dup")


def acumulo_abaixo_de_limites(serie, limites=None, faixa=0.05, min_qtd=5, razao_min=2.5):
    """Procura 'degraus' de valores logo abaixo de limites redondos — sinal de
    fracionamento para fugir de alçada de aprovação. Para cada limite L compara
    [L-5%, L) com a vizinhança. Sem `limites`, testa 1.000, 2.000, 2.500, 5.000,
    10.000, ... dentro da faixa dos dados. Devolve DataFrame com 'suspeito'."""
    v = _num(serie).dropna().abs()
    v = v[v > 0]
    colunas = ["limite", "logo_abaixo", "esperado", "razao", "suspeito"]
    if v.empty:
        return pd.DataFrame(columns=colunas)
    if limites is None:
        candidatos = sorted({m * 10 ** e for e in range(2, 10) for m in (1, 2, 2.5, 3, 5)})
        lo, hi = v.quantile(0.10), v.max()
        limites = [L for L in candidatos if lo < L * 0.85 and L * (1 - faixa) <= hi]
    linhas = []
    for L in limites:
        a = int(((v >= L * (1 - faixa)) & (v < L)).sum())
        b = int(((v >= L * (1 - 3 * faixa)) & (v < L * (1 - faixa))).sum())
        c = int(((v >= L) & (v < L * (1 + faixa))).sum())
        esperado = (b / 2 + c) / 2
        linhas.append({"limite": L, "logo_abaixo": a, "esperado": esperado,
                       "razao": a / max(esperado, 0.5)})
    out = pd.DataFrame(linhas, columns=colunas[:-1])
    out["suspeito"] = (out["logo_abaixo"] >= min_qtd) & (out["razao"] >= razao_min)
    return out


def fracionamento(df, col_valor, limite, col_entidade, col_data=None):
    """Grupos (entidade, dia) com 2+ valores abaixo do limite cuja soma passa
    do limite — o clássico 'quebrar um pagamento grande em vários pequenos'."""
    v = _num(df[col_valor])
    base = df[(v > 0) & (v < limite)]
    chaves = [base[col_entidade].astype(object).rename("entidade")]
    if col_data:
        chaves.append(converter_data(base[col_data]).dt.normalize().rename("dia"))
    g = _num(base[col_valor]).groupby(chaves).agg(qtd="size", soma="sum", maior="max")
    g = g[(g["qtd"] >= 2) & (g["soma"] >= limite)]
    return g.sort_values("soma", ascending=False)


def _pascoa(ano):
    """Domingo de Páscoa (algoritmo de Meeus/Jones/Butcher)."""
    a, b, c = ano % 19, ano // 100, ano % 100
    d, e = b // 4, b % 4
    f = (b + 8) // 25
    g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i, k = c // 4, c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    mes = (h + l - 7 * m + 114) // 31
    dia = (h + l - 7 * m + 114) % 31 + 1
    return date(ano, mes, dia)


def feriados_br(anos):
    """Feriados nacionais e pontos facultativos nacionais: {date: nome}.
    Inclui Carnaval, Sexta-feira Santa e Corpus Christi; Consciência Negra a partir de 2024."""
    if isinstance(anos, (int, np.integer)):
        anos = [anos]
    out = {}
    for ano in sorted({int(a) for a in anos if not _na(a) and 1900 <= int(a) <= 2200}):
        fixos = {(1, 1): "Confraternização Universal", (4, 21): "Tiradentes",
                 (5, 1): "Dia do Trabalho", (9, 7): "Independência",
                 (10, 12): "Nossa Senhora Aparecida", (11, 2): "Finados",
                 (11, 15): "Proclamação da República", (12, 25): "Natal"}
        if ano >= 2024:
            fixos[(11, 20)] = "Consciência Negra"
        for (mes, dia), nome in fixos.items():
            out[date(ano, mes, dia)] = nome
        p = _pascoa(ano)
        out[p - timedelta(days=48)] = "Carnaval (segunda)"
        out[p - timedelta(days=47)] = "Carnaval (terça)"
        out[p - timedelta(days=2)] = "Sexta-feira Santa"
        out[p + timedelta(days=60)] = "Corpus Christi"
    return dict(sorted(out.items()))


def marcar_calendario(datas, horas=None):
    """Para cada data: dia da semana, fim de semana, feriado (nome) e, se houver
    hora, madrugada (0h–5h59)."""
    d = converter_data(datas)
    out = pd.DataFrame(index=d.index)
    out["dia_semana"] = d.dt.dayofweek
    out["fim_de_semana"] = _bool(out["dia_semana"] >= 5)
    anos = d.dt.year.dropna().unique()
    fer = feriados_br(anos)
    mapa = pd.Series(list(fer.values()), index=pd.to_datetime(list(fer.keys())), dtype=object)
    out["feriado"] = d.dt.normalize().map(mapa) if len(mapa) else None
    if horas is None and d.notna().any():
        validas = d.dropna()
        if ((validas.dt.hour != 0) | (validas.dt.minute != 0)).any():
            horas = d.dt.hour
    if horas is not None:
        out["hora"] = horas
        out["madrugada"] = _bool(out["hora"].between(0, 5))
    return out


_PESOS_ALERTA = {"valor_atipico": 2, "duplicidade": 2, "fracionamento": 2, "abaixo_limite": 2,
                 "valor_redondo": 1, "madrugada": 1, "fim_de_semana": 1, "feriado": 1}
_NOMES_ALERTA = {"valor_atipico": "atípico", "duplicidade": "duplicidade",
                 "fracionamento": "fracionado", "abaixo_limite": "abaixo do limite",
                 "valor_redondo": "redondo", "madrugada": "madrugada",
                 "fim_de_semana": "fds", "feriado": "feriado"}


def sinalizar(df, col_valor, col_data=None, col_entidade=None, limite=None, horas=None, log=None):
    """Marca cada linha com alertas de auditoria e soma uma pontuação de risco.
    Colunas novas: alerta_* (bool), pontos_risco e motivos (texto).
    limite: número ou lista de limites de alçada (None = detecta degraus sozinho).
    log: z-score robusto na escala log (None = decide pela assimetria do valor).
    Ordene por 'pontos_risco' para saber o que conferir primeiro."""
    v = _num(df[col_valor])
    log = _assimetrica_positiva(v) if log is None else log
    flags = pd.DataFrame(index=df.index)
    flags["valor_atipico"] = zscore_robusto(v, log=log).abs() > 3.5
    flags["valor_redondo"] = (v.abs() >= 1000) & ((v.abs().round(2) % 1000) == 0)
    if col_data or col_entidade:
        dup = possiveis_duplicatas(df, col_valor, col_data, col_entidade)
        flags["duplicidade"] = df.index.isin(dup.index)
    if limite is None:
        acum = acumulo_abaixo_de_limites(v)
        limites = list(acum.loc[acum["suspeito"], "limite"])
    else:
        limites = list(limite) if isinstance(limite, (list, tuple, set)) else [limite]
    if limites:
        abaixo = pd.Series(False, index=df.index)
        for L in limites:
            abaixo |= (v.abs() >= L * 0.95) & (v.abs() < L)
        flags["abaixo_limite"] = abaixo
        if col_entidade:
            fr = pd.Series(False, index=df.index)
            for L in limites:
                grupos = fracionamento(df, col_valor, L, col_entidade, col_data)
                if len(grupos):
                    chave = df[col_entidade].astype(object)
                    if col_data:
                        dia = converter_data(df[col_data]).dt.normalize()
                        dentro = pd.MultiIndex.from_arrays([chave, dia]).isin(grupos.index)
                    else:
                        dentro = chave.isin(grupos.index).to_numpy()
                    fr |= pd.Series(dentro, index=df.index) & (v > 0) & (v < L)
            flags["fracionamento"] = fr
    if col_data:
        cal = marcar_calendario(df[col_data], horas)
        flags["fim_de_semana"] = cal["fim_de_semana"]
        flags["feriado"] = cal["feriado"].notna()
        if "madrugada" in cal:
            flags["madrugada"] = cal["madrugada"]
    flags = _bool(flags)
    out = df.copy()
    for c in flags.columns:
        out[f"alerta_{c}"] = flags[c]
    pesos = pd.Series({c: _PESOS_ALERTA.get(c, 1) for c in flags.columns})
    out["pontos_risco"] = (flags * pesos).sum(axis=1).astype(int)
    nomes = [_NOMES_ALERTA.get(c, c) for c in flags.columns]
    out["motivos"] = [", ".join(n for n, f in zip(nomes, linha) if f)
                      for linha in flags.itertuples(index=False, name=None)]
    return out


# ===========================================================================
# 9. CRUZAMENTO DE TABELAS — checar antes do merge
# ===========================================================================
def padronizar_chave(serie):
    """Chave de merge comparável: texto, sem espaços/acentos, maiúsculo, sem
    zeros à esquerda e sem o '.0' de números que viraram float ('00123' == '123.0')."""
    def um(v):
        if _na(v):
            return v
        if isinstance(v, (float, np.floating)) and float(v).is_integer():
            v = int(v)
        t = re.sub(r"\s+", " ", _sem_acento(str(v)).strip().upper())
        if re.fullmatch(r"[\d.\-/]+", t) and re.search(r"\d", t):
            t = re.sub(r"\D", "", t).lstrip("0") or "0"        # CPF/CNPJ formatado ou não
        return t
    return serie.map(um)


def _chave_combinada(df, cols, normalizar=False):
    partes = [padronizar_chave(df[c]) if normalizar else df[c] for c in cols]
    if len(partes) == 1:
        return partes[0]
    return pd.Series(list(zip(*partes)), index=df.index)


def checar_merge(esq, dir, chave=None, chave_esq=None, chave_dir=None, verbose=True):
    """Diagnóstico ANTES de juntar duas tabelas: chaves vazias e repetidas,
    tipos diferentes, % de casamento, quantas linhas o merge vai gerar e se
    padronizar o texto da chave ajudaria. Devolve um dicionário com tudo."""
    ke = chave_esq or chave
    kd = chave_dir or chave
    if ke is None or kd is None:
        raise ValueError("Informe chave= (mesmo nome nos dois lados) ou chave_esq= e chave_dir=.")
    ke = [ke] if isinstance(ke, str) else list(ke)
    kd = [kd] if isinstance(kd, str) else list(kd)
    ke = [resolver_coluna(esq, c) for c in ke]
    kd = [resolver_coluna(dir, c) for c in kd]
    ce, cd = _chave_combinada(esq, ke), _chave_combinada(dir, kd)
    lados = {}
    for nome, df, cols, k in (("esquerda", esq, ke, ce), ("direita", dir, kd, cd)):
        vc = k.value_counts(dropna=False)
        lados[nome] = {"linhas": len(df), "distintas": int(k.nunique()),
                       "vazias": int(df[cols].isna().any(axis=1).sum()),
                       "chaves_repetidas": int((vc > 1).sum()),
                       "linhas_repetidas": int(vc[vc > 1].sum()),
                       "tipos": [_nome_tipo(df[c]) for c in cols], "contagem": vc}
    tipos_diferentes = [f"{a} ({ta}) × {b} ({tb})" for a, b, ta, tb in
                        zip(ke, kd, lados["esquerda"]["tipos"], lados["direita"]["tipos"])
                        if ta != tb and not {ta, tb} <= {"inteiro", "decimal"}]
    casa_esq = ce.isin(set(cd.dropna())) if len(ke) == 1 else ce.isin(set(cd))
    casa_dir = cd.isin(set(ce.dropna())) if len(kd) == 1 else cd.isin(set(ce))
    taxa_esq = casa_esq.mean() * 100 if len(ce) else np.nan
    taxa_dir = casa_dir.mean() * 100 if len(cd) else np.nan
    ne, nd = _chave_combinada(esq, ke, True), _chave_combinada(dir, kd, True)
    taxa_norm = ne.isin(set(nd.dropna()) if len(ke) == 1 else set(nd)).mean() * 100 if len(ne) else np.nan
    vl, vr = lados["esquerda"]["contagem"], lados["direita"]["contagem"]
    comuns = vl.index.intersection(vr.index)
    inner = int((vl[comuns] * vr[comuns]).sum())
    left = inner + int(vl[~vl.index.isin(comuns)].sum())
    rel_e = "N" if lados["esquerda"]["chaves_repetidas"] else "1"
    rel_d = "N" if lados["direita"]["chaves_repetidas"] else "1"
    res = {
        "chave_esq": ke, "chave_dir": kd, "relacao": f"{rel_e}:{rel_d}",
        "pct_esq_com_par": taxa_esq, "pct_dir_com_par": taxa_dir,
        "pct_esq_com_par_padronizando": taxa_norm,
        "linhas_inner": inner, "linhas_left": left, "tipos_diferentes": tipos_diferentes,
        "sem_par_esq": ce[~casa_esq].dropna().drop_duplicates().head(10).tolist(),
        "sem_par_dir": cd[~casa_dir].dropna().drop_duplicates().head(10).tolist(),
        **{f"{lado}_{k}": v for lado, d in lados.items() for k, v in d.items() if k != "contagem"},
    }
    if verbose:
        subtitulo(f"Checagem do merge: {' + '.join(ke)} × {' + '.join(kd)}")
        linhas = [[rot] + [br(lados[l][k]) for l in ("esquerda", "direita")] for rot, k in
                  (("linhas", "linhas"), ("chaves distintas", "distintas"),
                   ("chaves vazias", "vazias"), ("chaves repetidas", "chaves_repetidas"))]
        linhas.append(["tipo da chave"] + [", ".join(lados[l]["tipos"]) for l in ("esquerda", "direita")])
        tabela(["", "esquerda", "direita"], linhas, ["l", "r", "r"], recuo=4)
        print()
        if tipos_diferentes:
            alerta(f"Tipos diferentes na chave: {'; '.join(tipos_diferentes)}. O pandas não casa "
                   "texto com número — converta um dos lados (ou use padronizar_chave).", 4)
        nota(f"{pct(taxa_esq)} das linhas da esquerda acham par na direita; "
             f"{pct(taxa_dir)} da direita acham par na esquerda. Relação {res['relacao']}.", 4)
        nota(f"Linhas após o merge: inner {br(inner)} · left {br(left)} "
             f"(a esquerda tem {br(len(esq))}).", 4)
        if left > len(esq):
            aviso(f"Chaves repetidas na direita vão DUPLICAR {br(left - len(esq))} linhas da "
                  "esquerda: somas depois do merge ficarão infladas.", 4)
        if not _na(taxa_norm) and not _na(taxa_esq) and taxa_norm > taxa_esq + 1:
            aviso(f"Padronizando a chave (espaços, caixa, acentos, zeros à esquerda) o casamento "
                  f"sobe para {pct(taxa_norm)}: aplique padronizar_chave() nos dois lados.", 4)
        if res["sem_par_esq"]:
            nota("Sem par (esquerda), ex: " + ", ".join(repr(v) for v in res["sem_par_esq"][:5]), 4)
    return res


# ===========================================================================
# 10. RELATÓRIO
# ===========================================================================
SECOES = ("visao", "qualidade", "ficha", "outliers", "extremos", "categorias", "entidades",
          "tempo", "correlacao", "fraude", "achados")


class Relatorio:
    """Resultado de analisar(): o df analisado, as colunas usadas, as tabelas
    de cada seção (rel.tabelas) e os achados em texto (rel.achados).
    Com antifraude, rel.alertas tem as linhas sinalizadas, da mais suspeita."""

    def __init__(self, df, info=None, top=5, debug=False):
        self.df = df
        self.info = info or {}
        self.top = top
        self.debug = debug
        self.col_valor = self.col_categoria = self.col_data = self.col_entidade = None
        self.limite = None
        self.escolhas = []
        self.cand_categoria, self.cand_data, self.cand_entidade = [], [], []
        self.pii = {}
        self.tabelas = {}
        self.achados = []
        self.alertas = None
        self.horas = None
        self._n = 0
        self.offset = self.info.get("linha_offset")

    def achado(self, texto, nivel="info"):
        self.achados.append((nivel, texto))

    def secao(self, texto):
        self._n += 1
        titulo(f"{self._n} · {texto}")

    def com_linha(self, sub):
        """Cópia com a coluna 'linha' (linha no arquivo original, se conhecida)."""
        out = sub.copy()
        nome = "linha" if self.offset is not None else "índice"
        if self.offset is not None and pd.api.types.is_integer_dtype(out.index):
            out.insert(0, nome, out.index + self.offset)
        else:
            out.insert(0, nome, out.index)
        return out

    def chaves(self):
        """Colunas que importam primeiro numa tabela de linhas."""
        return [c for c in ("linha", "índice", self.col_valor, self.col_entidade, self.col_data,
                            self.col_categoria) if c]

    def __repr__(self):
        return (f"<Relatorio: {len(self.df)} linhas · valor={self.col_valor!r} · "
                f"categoria={self.col_categoria!r} · data={self.col_data!r} · "
                f"{len(self.achados)} achados · tabelas: {', '.join(self.tabelas)}>")


def _rodar(rel, nome, funcao):
    try:
        funcao(rel)
    except Exception as erro:
        print()
        aviso(f"A seção '{nome}' não pôde ser concluída: {type(erro).__name__}: {erro}")
        if rel.debug:
            traceback.print_exc()


def _mostrar_linhas(rel, sub, max_linhas=10, extras=None, recuo=4, max_cel=30):
    tab = rel.com_linha(sub)
    prioridade = tab.columns[:1].tolist() + (extras or []) + rel.chaves()
    mostrar(tab, indice=False, recuo=recuo, max_linhas=max_linhas, prioridade=prioridade,
            casas=_casas(rel.df[rel.col_valor]) if rel.col_valor else 2, max_cel=max_cel)


def _escolher_colunas(rel, col_valor, col_categoria, col_data, col_entidade, padroes):
    """Resolve as colunas pedidas (erro se não existirem) e, quando faltam,
    usa os padrões do CONFIG (sem erro) ou escolhe sozinho."""
    df = rel.df
    padroes = padroes or {}
    rel.pii = detectar_pii(df)
    mudou = False

    def suave(nome):
        if not nome:
            return None
        try:
            return resolver_coluna(df, nome)
        except KeyError:
            rel.escolhas.append(("aviso", f"Coluna padrão '{nome}' (CONFIG) não existe nesta base; "
                                          "ignorada."))
            return None

    # valor
    if col_valor is not None:
        c = resolver_coluna(df, col_valor)
        modo = "pedida (-v)"
    else:
        c = suave(padroes.get("valor"))
        modo = "do CONFIG"
        if c is not None and not _eh_numerica(df[c]):
            c = None
        if c is None:
            c = escolher_coluna_valor(df)
            modo = "escolhida automaticamente (troque com -v)"
    if c is not None and not _eh_numerica(df[c]):
        conv = converter_numero(df[c])
        if conv.isna().all():
            conv = converter_numero(df[c], "us")
        taxa = conv.notna().sum() / max(df[c].notna().sum(), 1)
        if taxa < 0.8:
            exemplos = ", ".join(repr(cortar(v, 15)) for v in df[c].dropna().head(5))
            raise ValueError(f"A coluna '{c}' não é numérica (tipo: {_nome_tipo(df[c])}). "
                             f"Exemplos: {exemplos}")
        df = df.copy()
        df[c] = conv
        mudou = True
        modo += f"; convertida para número ({pct((1 - taxa) * 100)} viraram vazio)"
    rel.col_valor = c

    # data
    candidatas = colunas_data(df)
    if col_data:
        c = resolver_coluna(df, col_data)
        modo_d = "pedida (-d)"
    else:
        c = suave(padroes.get("data"))
        modo_d = "do CONFIG"
        if c is None:
            eventos = [x for x in candidatas if not set(_tokens(x)) & _DATA_NAO_EVENTO]
            if len(eventos) == 1:
                c, modo_d = eventos[0], "única coluna de data (troque com -d)"
    if c is not None and not pd.api.types.is_datetime64_any_dtype(df[c]):
        conv = converter_data(df[c])
        if conv.notna().sum() == 0:
            raise ValueError(f"A coluna '{c}' não tem datas reconhecíveis. Exemplos: "
                             + ", ".join(repr(cortar(v, 15)) for v in df[c].dropna().head(5)))
        if not mudou:
            df = df.copy()
            mudou = True
        df[c] = conv
    rel.col_data = c
    rel.cand_data = [x for x in candidatas if x != c]

    # categoria e entidade
    rel.col_categoria = resolver_coluna(df, col_categoria) if col_categoria else suave(padroes.get("categoria"))
    rel.col_entidade = resolver_coluna(df, col_entidade) if col_entidade else suave(padroes.get("entidade"))
    usadas = {rel.col_valor, rel.col_data, rel.col_categoria, rel.col_entidade}
    rel.cand_categoria = [x for x in colunas_categoricas(df) if x not in usadas]
    rel.cand_entidade = [x for x in colunas_entidade(df, rel.pii) if x not in usadas]
    rel.df = df

    def linha(papel, col, modo_txt, cands, opcao):
        if col is not None:
            return [papel, str(col), modo_txt]
        if cands:
            return [papel, "—", f"candidatas: {', '.join(map(str, cands[:5]))} (use {opcao})"]
        return [papel, "—", "nenhuma candidata"]

    rel.escolhas_tabela = [
        linha("valor", rel.col_valor, modo, [], "-v"),
        linha("categoria", rel.col_categoria, "pedida (-c)" if col_categoria else "do CONFIG",
              rel.cand_categoria, "-c"),
        linha("data", rel.col_data, modo_d, rel.cand_data, "-d"),
        linha("entidade", rel.col_entidade, "pedida (-e)" if col_entidade else "do CONFIG",
              rel.cand_entidade, "-e"),
    ]
    if rel.col_valor is None:
        rel.escolhas.append(("aviso", "Nenhuma coluna numérica: as análises de valor serão puladas; "
                                      "tempo e categorias usam contagem de linhas."))
    if rel.col_data:
        rel.horas = _serie_horas(df, rel.col_data)


# --- seções ------------------------------------------------------------------
def _amostra(df, n=3, offset=None, recuo=2):
    sub = df.head(n)
    cab = [str(c) for c in sub.columns]
    cels = [[_celula(v, max_cel=24) for v in row] for row in sub.itertuples(index=False, name=None)]
    larg = sum(max([len(c)] + [len(r[i]) for r in cels]) for i, c in enumerate(cab)) + 2 * (len(cab) - 1)
    if larg <= _SAIDA["largura"] - recuo:
        mostrar(sub, indice=False, recuo=recuo)
        return
    rot = [f"linha {i + offset}" if offset is not None and isinstance(i, (int, np.integer))
           else f"#{i}" for i in sub.index]
    linhas = [[cortar(c, 26)] + [cels[j][i] for j in range(len(cels))] for i, c in enumerate(cab)]
    tabela(["coluna"] + rot, linhas, ["l"] * (len(rot) + 1), recuo)


def dicionario_dados(df, pii=None):
    """Uma linha por coluna: tipo, papel (que dado é esse?), nulos, distintos, exemplo."""
    pii = detectar_pii(df) if pii is None else pii
    linhas = []
    for col in df.columns:
        s = df[col]
        nulos = int(s.isna().sum())
        preenchidos = s.dropna()
        linhas.append({
            "coluna": col, "tipo": _nome_tipo(s), "papel": _papel(df, col, pii), "nulos": nulos,
            "% nulos": nulos / len(s) * 100 if len(s) else 0.0, "distintos": _nunique(s),
            "exemplo": _celula(preenchidos.iloc[0], max_cel=24) if len(preenchidos) else "-",
        })
    return pd.DataFrame(linhas, columns=["coluna", "tipo", "papel", "nulos", "% nulos",
                                         "distintos", "exemplo"])


def _tabela_dicionario(d):
    tabela(list(d.columns),
           [[cortar(r["coluna"], 28), r["tipo"], r["papel"], br(r["nulos"]), br(r["% nulos"], 1),
             br(r["distintos"]), r["exemplo"]] for _, r in d.iterrows()],
           ["l", "l", "l", "r", "r", "r", "l"])


def visao_geral(df, linhas_amostra=3):
    """Formato, dicionário de dados e amostra, impressos."""
    nota(f"{br(len(df))} linhas × {len(df.columns)} colunas")
    subtitulo("Colunas — que dado é esse?")
    d = dicionario_dados(df)
    _tabela_dicionario(d)
    subtitulo(f"Amostra ({min(linhas_amostra, len(df))} linhas)")
    _amostra(df, linhas_amostra)
    return d


def _secao_visao(rel):
    df = rel.df
    rel.secao("VISÃO GERAL")
    memoria = df.memory_usage(deep=len(df) <= 2_000_000).sum()
    nota(f"{br(len(df))} linhas × {len(df.columns)} colunas · {_bytes(memoria)} na memória")
    subtitulo("Colunas — que dado é esse?")
    d = dicionario_dados(df, rel.pii)
    rel.tabelas["Dicionário"] = d
    _tabela_dicionario(d)
    subtitulo("Amostra")
    _amostra(df, 3, rel.offset)
    subtitulo("Colunas usadas na análise")
    tabela(["papel", "coluna", "como"], rel.escolhas_tabela, recuo=4)
    for nivel, texto in rel.escolhas:
        _status(nivel, texto, 4)


def _secao_qualidade(rel):
    df = rel.df
    rel.secao("QUALIDADE DOS DADOS")
    subtitulo("Vazios")
    nulos = diagnostico_nulos(df)
    if nulos.empty:
        _status("ok", "Nenhum valor vazio.", 4)
    else:
        tabela(["coluna", "vazios", "%", ""],
               [[cortar(c, 30), br(int(r["nulos"])), br(r["pct"], 1), barra(r["pct"], 100)]
                for c, r in nulos.iterrows()], ["l", "r", "r", "l"], recuo=4)
        nota("Antes de apagar ou preencher, descubra POR QUE está vazio.", 4)
        if rel.col_valor in nulos.index and nulos.loc[rel.col_valor, "pct"] >= 1:
            rel.achado(f"{pct(nulos.loc[rel.col_valor, 'pct'])} dos registros estão sem "
                       f"{rel.col_valor}.", "aviso")
    subtitulo("Duplicatas")
    try:
        dup = int(df.duplicated().sum())
    except TypeError:
        dup = int(df.astype(str).duplicated().sum())
    if dup == 0:
        _status("ok", "Nenhuma linha duplicada.", 4)
    else:
        aviso(f"{br(dup)} linhas são cópia exata de outra ({pct(dup / len(df) * 100)} da base).", 4)
        _mostrar_linhas(rel, duplicatas(df).head(rel.top * 2), max_linhas=rel.top * 2)
        nota("Para remover: df = df.drop_duplicates()", 4)
        rel.achado(f"{br(dup)} linhas duplicadas (cópia exata) — somas ficam infladas se não "
                   "forem removidas.", "aviso")
    subtitulo("Checagens automáticas")
    q = checar_qualidade(df, rel.col_valor, rel.pii)
    rel.tabelas["Qualidade"] = q
    feitas = ["colunas vazias/constantes", "IDs repetidos", "datas futuras/implausíveis",
              "grafias diferentes", "espaços sobrando", "números como texto", "CPF/CNPJ inválidos",
              "dados pessoais"]
    if q.empty:
        _status("ok", "Nada estranho em: " + ", ".join(feitas) + ".", 4)
    ordem = {"alerta": 0, "aviso": 1, "info": 2}
    for _, r in q.sort_values("nivel", key=lambda s: s.map(ordem)).iterrows():
        _status(r["nivel"], f"{r['coluna']} — {r['checagem']}: {r['detalhe']}", 4)
        if r["nivel"] in ("alerta", "aviso"):
            rel.achado(f"{r['coluna']}: {r['checagem']} — {r['detalhe']}", r["nivel"])


def _secao_ficha(rel):
    cv = rel.col_valor
    s = _num(rel.df[cv])
    rel.secao(f"FICHA DA VARIÁVEL: {cv}")
    if s.notna().sum() < 2:
        nota("Menos de 2 valores preenchidos: não dá para descrever.")
        return
    r = resumo_numerico(s)
    c = _casas(s)
    rel.tabelas["Ficha"] = r.rename("valor").rename_axis("indicador").reset_index()
    subtitulo("1. Que dado é esse?")
    inteiro = bool((s.dropna() % 1 == 0).all())
    par("tipo", "numérica " + ("discreta (inteiros)" if inteiro else "contínua (decimais)"))
    par("preenchidos", f"{br(r['contagem'])} de {br(len(s))}"
        + (f"  ({br(r['nulos'])} vazios)" if r["nulos"] else ""))
    par("distintos", br(_nunique(s)))
    par("soma", br(r["soma"], c))
    if r["negativos"] or r["zeros"]:
        par("negativos / zeros", f"{br(r['negativos'])} / {br(r['zeros'])}")
    subtitulo("2. Qual é o valor típico?")
    par("média", br(r["media"], c))
    par("mediana", br(r["mediana"], c))
    par("moda", f"{br(r['moda'], c)}  (aparece {br(r['moda_freq'])} vezes)" if r["moda_freq"]
        else "nenhum valor se repete")
    subtitulo("3. Quanto varia?")
    par("desvio padrão", br(r["desvio_padrao"], c))
    par("coef. de variação", pct(r["coef_variacao_pct"]) + ("  (não confiável: há negativos)"
                                                           if r["negativos"] else ""))
    par("amplitude", br(r["amplitude"], c))
    subtitulo("4. Onde cada valor está?")
    tabela(["mínimo", "Q1 (25%)", "mediana", "Q3 (75%)", "máximo"],
           [[br(r["min"], c), br(r["q25"], c), br(r["mediana"], c), br(r["q75"], c), br(r["max"], c)]],
           ["r"] * 5, recuo=4)
    h = histograma(s)
    if len(h) > 1:
        print()
        maximo = h["qtd"].max()
        tabela(["faixa", "qtd", "%", ""],
               [[r_["faixa"] + (" ⟵ fora" if r_["fora"] else ""), br(int(r_["qtd"])),
                 br(r_["pct"], 1), barra(r_["qtd"], maximo)] for _, r_ in h.iterrows()],
               ["l", "r", "r", "l"], recuo=4)
        rel.tabelas["Histograma"] = h
    subtitulo("Leitura em uma frase")
    frase = leitura_em_frase(s, casas=c)
    _escrever(frase, 4)
    rel.achado(frase.split(". ")[0].rstrip(".") + ".")
    if abs(r["media"] - r["mediana"]) > 0.10 * abs(r["mediana"] or 1) and r["mediana"]:
        rel.achado(f"Média ({br(r['media'], c)}) e mediana ({br(r['mediana'], c)}) de {cv} "
                   "diferem bastante: a distribuição é assimétrica, fale do típico pela mediana.")


def _secao_outliers(rel):
    df, cv = rel.df, rel.col_valor
    s = _num(df[cv])
    c = _casas(s)
    rel.secao(f"OUTLIERS — regra do IQR ({cv})")
    inf, sup = limites_iqr(s)
    nota(f"Faixa normal (Q1 − 1,5·IQR até Q3 + 1,5·IQR): {br(inf, c)} a {br(sup, c)}")
    outs = outliers_iqr(df, cv)
    if outs.empty:
        _status("ok", "Nenhum valor fora da faixa.", 2)
        return
    v = _num(outs[cv])
    e_inf, e_sup = limites_iqr(s, 3)
    extremos = int(((s < e_inf) | (s > e_sup)).sum())
    if limites_iqr(s, 0)[0] == limites_iqr(s, 0)[1]:
        aviso("IQR = 0: mais da metade dos valores é igual, então qualquer valor diferente "
              "vira 'outlier'. Prefira o z-score robusto (seção antifraude, -f).")
    aviso(f"{br(len(outs))} valores fora da faixa ({pct(len(outs) / s.notna().sum() * 100)}): "
          f"{br(int((v > sup).sum()))} acima e {br(int((v < inf).sum()))} abaixo; "
          f"{br(extremos)} são extremos (além de 3·IQR).")
    ordem = (v - s.median()).abs().sort_values(ascending=False).index
    print()
    _mostrar_linhas(rel, outs.loc[ordem], max_linhas=10)
    nota(f"Lista completa: outliers_iqr(df, {cv!r})", 4)
    rel.tabelas["Outliers"] = rel.com_linha(outs.loc[ordem])
    rel.achado(f"{br(len(outs))} valores de {cv} fogem do padrão (acima de {br(sup, c)}"
               + (f" ou abaixo de {br(inf, c)}" if (v < inf).any() else "")
               + f"); {br(extremos)} são extremos.", "aviso" if extremos else "info")


def _secao_extremos(rel):
    df, cv = rel.df, rel.col_valor
    rel.secao(f"MAIORES E MENORES ({cv})")
    validos = df[_num(df[cv]).notna()]
    subtitulo(f"{rel.top} maiores")
    _mostrar_linhas(rel, validos.loc[_num(validos[cv]).nlargest(rel.top).index], rel.top)
    subtitulo(f"{rel.top} menores")
    _mostrar_linhas(rel, validos.loc[_num(validos[cv]).nsmallest(rel.top).index], rel.top)


def _tabela_grupo(rel, chave, limite=15, recuo=4):
    cv = rel.col_valor
    g = resumo_por_grupo(rel.df, chave, cv, limite)
    c = _casas(rel.df[cv]) if cv else 0
    if cv:
        maximo = g["total"].abs().max()
        linhas = [[cortar(k, 26), br(int(r["qtd"])), br(r["total"], c), br(r["pct"], 1),
                   br(r["pct_acum"], 1), r["classe"], br(r["media"], c), barra(r["total"], maximo)]
                  for k, r in g.iterrows()]
        tabela([str(chave), "qtd", "total", "%", "% acum", "ABC", "média", ""], linhas,
               ["l", "r", "r", "r", "r", "l", "r", "l"], recuo)
    else:
        maximo = g["qtd"].max()
        linhas = [[cortar(k, 26), br(int(r["qtd"])), br(r["pct"], 1), br(r["pct_acum"], 1),
                   r["classe"], barra(r["qtd"], maximo)] for k, r in g.iterrows()]
        tabela([str(chave), "qtd", "%", "% acum", "ABC", ""], linhas,
               ["l", "r", "r", "r", "l", "l"], recuo)
    return curva_abc(rel.df, chave, cv)


def _leitura_grupos(rel, g, unidade="categorias"):
    cv = rel.col_valor
    c = _casas(rel.df[cv]) if cv else 0
    base = "total" if cv else "qtd"
    lider, r0 = g.index[0], g.iloc[0]
    frases = [f"'{lider}' lidera com {pct(r0['pct'])} do " + ("valor total" if cv else "total de registros")
              + (f" ({br(r0['total'], c)} em {br(int(r0['qtd']))} registros)." if cv else ".")]
    n_a = int((g["classe"] == "A").sum())
    if len(g) > 3 and (g[base] >= 0).all():
        frases.append(f"{br(n_a)} de {br(len(g))} {unidade} ({pct(n_a / len(g) * 100, 0)}) "
                      f"{'concentra' if n_a == 1 else 'concentram'} ~80% (classe A).")
    if cv and len(g) > 1:
        media_geral = _num(rel.df[cv]).mean()
        relevantes = g[g["qtd"] >= max(5, 0.01 * len(rel.df))]
        if len(relevantes) and media_geral:
            top_m = relevantes["media"].idxmax()
            razao = relevantes.loc[top_m, "media"] / media_geral
            if razao >= 1.5:
                frases.append(f"'{top_m}' tem o maior valor médio ({br(relevantes.loc[top_m, 'media'], c)}), "
                              f"{br(razao, 1)}× a média geral.")
    if (g[base] < 0).any():
        frases.append("Há grupos com total negativo: os % são sobre o total líquido.")
    return " ".join(frases)


def _secao_categorias(rel):
    if rel.col_categoria:
        cat = rel.col_categoria
        rel.secao(f"{(rel.col_valor or 'REGISTROS').upper()} POR {str(cat).upper()}")
        g = _tabela_grupo(rel, cat)
        rel.tabelas[f"Por {cat}"] = g.reset_index()
        divididas = grafias_inconsistentes(rel.df[cat]) if _eh_texto(rel.df[cat]) else []
        if len(divididas):
            print()
            funcao = "padronizar_uf" if _eh_coluna_uf(cat) else "padronizar_texto"
            aviso(f"A mesma categoria aparece escrita de jeitos diferentes "
                  f"({' / '.join(map(repr, divididas['variantes'].iloc[0][:3]))}): os totais acima "
                  f"estão divididos. Padronize antes de concluir: df[{cat!r}] = "
                  f"{funcao}(df[{cat!r}]).", 4)
        print()
        frase = _leitura_grupos(rel, g)
        _escrever(frase, 4)
        rel.achado(frase.split(". ")[0].rstrip(".") + ".")
        return
    rel.secao("DISTRIBUIÇÃO DAS CATEGORIAS")
    nota("Nenhuma categoria escolhida (-c); resumo das candidatas:")
    for cat in rel.cand_categoria[:3]:
        subtitulo(str(cat))
        _tabela_grupo(rel, cat, limite=8)


def _secao_entidades(rel):
    ent, cv = rel.col_entidade, rel.col_valor
    rel.secao(f"ENTIDADES: {ent}")
    g = curva_abc(rel.df, ent, cv)
    rel.tabelas[f"Por {ent}"] = g.reset_index()
    base = "total" if cv else "qtd"
    unicas = int((g["qtd"] == 1).sum())
    nota(f"{br(len(g))} valores distintos de {ent}; {plural(unicas, 'aparece', 'aparecem')} uma única vez.")
    if (g[base] >= 0).all() and g[base].sum() > 0:
        top1 = g["pct"].iloc[0]
        n10 = max(1, int(round(len(g) * 0.1)))
        top10 = g["pct"].iloc[:n10].sum()
        nota(f"O maior responde por {pct(top1)} do total; os 10% maiores ({br(n10)}) por {pct(top10)}.")
    subtitulo(f"Top {min(10, len(g))}")
    _tabela_grupo(rel, ent, limite=10)
    print()
    frase = _leitura_grupos(rel, g, unidade="valores de " + str(ent))
    _escrever(frase, 4)
    rel.achado(". ".join(frase.split(". ")[:2]).rstrip(".") + ".")


def _tendencia_robusta(y):
    """Inclinação de Theil-Sen: mediana das inclinações entre todos os pares de
    pontos. Um pico isolado não puxa a tendência, ao contrário da regressão comum."""
    y = np.asarray(y, dtype=float)
    i, j = np.triu_indices(len(y), k=1)
    return float(np.median((y[j] - y[i]) / (j - i))) if len(i) else 0.0


def _datas_fora_da_curva(d, col):
    """Datas que distorceriam a série: no futuro (exceto colunas tipo vencimento),
    antes de 1900, ou isoladas longe do grosso dos dados (< 0,2% dos registros)."""
    validas = d.dropna()
    fora = pd.Series(False, index=d.index)
    if validas.empty:
        return fora
    fora |= d < pd.Timestamp("1900-01-01")
    if not set(_tokens(col)) & _DATA_NAO_EVENTO:
        fora |= d > pd.Timestamp.today().normalize() + pd.Timedelta(days=1)
    if len(validas) >= 500:
        lo, hi = validas.quantile(0.002), validas.quantile(0.998)
        margem = max(pd.Timedelta(days=31), (hi - lo) * 0.05)
        isoladas = (d < lo - margem) | (d > hi + margem)
        if isoladas.sum() <= 0.002 * len(validas) + 1:
            fora |= isoladas
    return _bool(fora)


def _secao_tempo(rel):
    df, cd, cv = rel.df, rel.col_data, rel.col_valor
    rel.secao(f"{(cv or 'REGISTROS').upper()} AO LONGO DO TEMPO ({cd})")
    fora = _datas_fora_da_curva(df[cd], cd)
    if fora.any():
        exemplos = ", ".join(data_br(x) for x in df.loc[fora, cd].drop_duplicates().head(3))
        aviso(f"{plural(int(fora.sum()), 'registro')} com data no futuro ou isolada longe das "
              f"demais ({exemplos}) ficaram fora desta seção, para não distorcer a série.")
        df = df[~fora]
    d = df[cd]
    validas = d.dropna()
    if validas.empty:
        nota("Nenhuma data válida.")
        return
    ini, fim = validas.min(), validas.max()
    t, freq = resumo_temporal(df, cd, cv)
    nome = {"D": "dia", "W": "semana", "MS": "mês", "QS": "trimestre", "YS": "ano"}[freq]
    nota(f"De {data_br(ini)} a {data_br(fim)} ({br((fim - ini).days + 1)} dias) · agrupado por {nome}"
         + (f" · {br(int(d.isna().sum()))} registros sem data" if d.isna().any() else ""))
    medida = "total" if cv else "qtd"
    c = _casas(df[cv]) if cv else 0
    mostrar_t = t.tail(24)
    maximo = mostrar_t[medida].abs().max()
    linhas = []
    for _, r in mostrar_t.iterrows():
        var = "parcial" if r["parcial"] else pct_sinal(r["var_pct"])
        rot = r["rotulo"] + ("*" if r["parcial"] else "")
        if cv:
            linhas.append([rot, br(int(r["qtd"])), br(r["total"], c), var, barra(r[medida], maximo)])
        else:
            linhas.append([rot, br(int(r["qtd"])), var, barra(r[medida], maximo)])
    print()
    if cv:
        tabela([nome, "qtd", "total", "var.", ""], linhas, ["l", "r", "r", "r", "l"], recuo=4)
    else:
        tabela([nome, "qtd", "var.", ""], linhas, ["l", "r", "r", "l"], recuo=4)
    if len(t) > 24:
        nota(f"(mostrando os últimos 24 de {len(t)} períodos)", 4)
    if t["parcial"].any():
        nota("* período incompleto: os dados começam/terminam no meio dele, a variação não é comparável.", 4)
    rel.tabelas["Tempo"] = t.drop(columns=["periodo"])

    completos = t[~t["parcial"]]
    frases = []
    if len(completos) >= 2:
        i_max, i_min = completos[medida].idxmax(), completos[medida].idxmin()
        media_p = completos[medida].mean()
        unid = "" if cv else " registros"
        frases.append(f"Maior {nome}: {completos.loc[i_max, 'rotulo']} "
                      f"({br(completos.loc[i_max, medida], c)}{unid}, "
                      f"{pct_sinal((completos.loc[i_max, medida] / media_p - 1) * 100)} vs a média). "
                      f"Menor: {completos.loc[i_min, 'rotulo']} ({br(completos.loc[i_min, medida], c)}{unid}).")
        if cv and completos.loc[i_max, "total"] > 0:
            ini_p = completos.loc[i_max, "periodo"]
            fim_p = ini_p + pd.tseries.frequencies.to_offset(_FREQS[freq][0])
            no_periodo = _num(df.loc[(d >= ini_p) & (d < fim_p), cv])
            if len(no_periodo) > 1:
                parte = no_periodo.max() / completos.loc[i_max, "total"]
                if parte >= 0.3:
                    frases.append(f"Atenção: um único registro de {br(no_periodo.max(), c)} responde "
                                  f"por {pct(parte * 100, 0)} de {completos.loc[i_max, 'rotulo']}.")
        ult, ant = completos.iloc[-1], completos.iloc[-2]
        if ant[medida]:
            var = (ult[medida] / ant[medida] - 1) * 100
            fem = nome == "semana"
            frases.append(f"{'Última' if fem else 'Último'} {nome} {'completa' if fem else 'completo'} "
                          f"({ult['rotulo']}): {pct_sinal(var)} sobre {'a anterior' if fem else 'o anterior'}.")
    if len(completos) >= 4 and completos[medida].median():
        y = completos[medida].to_numpy(dtype=float)
        inclinacao = _tendencia_robusta(y) / np.median(y) * 100
        tendencia = ("de alta" if inclinacao > 2 else "de queda" if inclinacao < -2 else "estável")
        frases.append(f"Tendência {tendencia} ({pct_sinal(inclinacao)} por {nome}, pela mediana das "
                      f"inclinações — um período fora da curva não muda a conclusão).")
    vazios = completos[completos["qtd"] == 0]
    if len(vazios):
        frases.append(f"Sem nenhum registro em: {', '.join(vazios['rotulo'].head(5))}"
                      + ("…" if len(vazios) > 5 else "") + " — buraco na extração?")
        rel.achado(f"{plural(len(vazios), nome, _PLURAIS.get(nome))} sem nenhum registro "
                   f"({', '.join(vazios['rotulo'].head(3))}).",
                   "aviso")
    if frases:
        subtitulo("Leitura")
        _escrever(" ".join(frases), 4)
        rel.achado(" ".join(f for f in frases if not f.startswith("Sem nenhum")))

    subtitulo("Por dia da semana")
    dow = validas.dt.dayofweek
    v = _num(df.loc[validas.index, cv]) if cv else None
    linhas, maximo = [], dow.value_counts().max()
    tot_v = v.sum() if v is not None else None
    for i, nome_dia in enumerate(_DIAS):
        m = dow == i
        q = int(m.sum())
        linha = [nome_dia, br(q), pct(q / len(dow) * 100)]
        if v is not None:
            linha += [br(v[m].sum(), c), pct(v[m].sum() / tot_v * 100) if tot_v else "-"]
        linhas.append(linha + [barra(q, maximo)])
    cab = ["dia", "qtd", "%"] + (["total", "% total"] if v is not None else []) + [""]
    tabela(cab, linhas, ["l", "r", "r"] + (["r", "r"] if v is not None else []) + ["l"], recuo=4)
    if rel.horas is not None:
        h = rel.horas.loc[validas.index].dropna()
        if len(h):
            subtitulo("Por faixa de horário")
            faixas = [("madrugada (0h–5h)", 0, 5), ("manhã (6h–11h)", 6, 11),
                      ("tarde (12h–17h)", 12, 17), ("noite (18h–23h)", 18, 23)]
            qtds = [int(h.between(a, b).sum()) for _, a, b in faixas]
            tabela(["faixa", "qtd", "%", ""],
                   [[n, br(q), pct(q / len(h) * 100), barra(q, max(qtds))]
                    for (n, _, _), q in zip(faixas, qtds)], ["l", "r", "r", "l"], recuo=4)
            pico = int(h.value_counts().idxmax())
            nota(f"Hora de pico: {pico}h ({pct(h.eq(pico).mean() * 100)} dos registros).", 4)


def _secao_correlacao(rel):
    df = rel.df
    nums = [c for c in df.columns if _eh_numerica(df[c]) and not parece_id(df, c)
            and _nunique(df[c]) >= 3]
    if len(nums) < 2:
        return
    nums = sorted(nums, key=lambda c: df[c].isna().sum())[:15]
    rel.secao("CORRELAÇÃO ENTRE COLUNAS NUMÉRICAS")
    nota("Perto de +1 ou −1 = andam juntas; perto de 0 = pouca relação. Correlação não é causa.")
    m = correlacao(df, nums)
    rel.tabelas["Correlação"] = m.reset_index().rename(columns={"index": "coluna"})
    if len(nums) <= 6:
        print()
        mostrar(m, casas=2, recuo=4, max_linhas=None)
    pares = pares_correlacao(df, nums, minimo=0.3)
    subtitulo("Pares mais relacionados")
    if pares.empty:
        nota("Nenhum par com correlação relevante (|r| ≥ 0,3).", 4)
        return
    tabela(["coluna A", "coluna B", "r", "força", ""],
           [[cortar(r["coluna_a"], 24), cortar(r["coluna_b"], 24), br(r["r"], 2),
             f"{r['forca']}, {r['direcao']}", barra(r["r"], 1)] for _, r in pares.head(8).iterrows()],
           ["l", "l", "r", "l", "l"], recuo=4)
    forte = pares[pares["forca"] == "forte"]
    if len(forte):
        r = forte.iloc[0]
        rel.achado(f"{r['coluna_a']} e {r['coluna_b']} têm correlação forte ({br(r['r'], 2)}).")


def _secao_fraude(rel):
    df, cv, cd, ce = rel.df, rel.col_valor, rel.col_data, rel.col_entidade
    s = _num(df[cv])
    c = _casas(s)
    rel.secao(f"ANTIFRAUDE ({cv})")
    nota("Sinais, não provas: cada alerta é um convite para conferir, não uma acusação.")

    subtitulo("Lei de Benford — primeiro dígito")
    b = benford(s)
    n = b.attrs["n"]
    if n < 50:
        nota(f"Só {n} valores diferentes de zero: poucos para o teste (mínimo 50).", 4)
    else:
        if n < 300:
            aviso("Menos de 300 valores: o teste serve só como indicação.", 4)
        if b.attrs["ordens"] < 2:
            aviso(f"Os valores variam só {br(b.attrs['ordens'], 1)} ordens de grandeza (o ideal é "
                  "≥ 2): Benford pode não se aplicar a esta coluna (preços tabelados, limites...).", 4)
        maximo = max(b["observado"].max(), b["esperado"].max())
        tabela(["dígito", "esperado", "observado", "dif. (pp)", "", ""],
               [[str(int(r["digito"])), pct(r["esperado"] * 100), pct(r["observado"] * 100),
                 ("+" if r["diferenca"] > 0 else "") + br(r["diferenca"] * 100, 1),
                 "!" if r["z"] > 1.96 else "", _barra_ref(r["observado"], r["esperado"], maximo)]
                for _, r in b.iterrows()], ["r", "r", "r", "r", "l", "l"], recuo=4)
        nota("│ = esperado por Benford · ! = diferença estatisticamente significativa (z > 1,96)", 4)
        conf = b.attrs["conformidade"]
        p = b.attrs["p_valor"]
        _status("alerta" if conf == "não conforme" else "aviso" if conf == "marginal" else "ok",
                f"MAD = {br(b.attrs['mad'], 4)} → conformidade {conf} (critério de Nigrini). "
                f"Qui² = {br(b.attrs['chi2'], 1)}, p {'= ' + br(p, 3) if p >= 0.001 else '< 0,001'}.", 4)
        nove = b[b["digito"] == 9].iloc[0]
        if nove["z"] > 1.96 and nove["diferenca"] > 0:
            nota("Excesso de valores começando com 9 costuma aparecer quando há valores logo "
                 "abaixo de limites (9.xxx < 10.000) — veja a parte de fracionamento.", 4)
        rel.tabelas["Benford"] = b
        b2 = benford(s, 2)
        picos = b2[(b2["z"] > 2.58) & (b2["diferenca"] > 0)].nlargest(5, "diferenca")
        if len(picos):
            txt = ", ".join(f"{int(r['digito'])} (+{br(r['diferenca'] * 100, 2)} pp)" for _, r in picos.iterrows())
            nota(f"Dois primeiros dígitos acima do esperado: {txt}. Filtre os valores que começam "
                 "assim — é onde costuma estar o problema.", 4)
        aplicavel = n >= 300 and b.attrs["ordens"] >= 2
        if conf in ("não conforme", "marginal") and aplicavel:
            rel.achado(f"Benford: {cv} tem conformidade {conf} (MAD {br(b.attrs['mad'], 4)})"
                       + (f"; dígitos iniciais em excesso: {', '.join(str(int(x)) for x in picos['digito'])}"
                          if len(picos) else "") + ".", "alerta" if conf == "não conforme" else "aviso")

    subtitulo("Valores atípicos — z-score robusto (MAD)")
    usa_log = _assimetrica_positiva(s)
    z = zscore_robusto(s, log=usa_log)
    atip = z.abs() > 3.5
    nota(f"{plural(int(atip.sum()), 'valor', 'valores')} com |z| > 3,5 ({pct(atip.mean() * 100)}). "
         "Usa mediana e MAD, então não é enganado pelos próprios outliers como o z-score comum"
         + ("; calculado na escala log, porque valores em dinheiro têm cauda longa." if usa_log else "."), 4)
    if atip.any():
        sub = df[atip].assign(z_robusto=z[atip])
        sub = sub.loc[sub["z_robusto"].abs().sort_values(ascending=False).index]
        _mostrar_linhas(rel, sub, max_linhas=rel.top, extras=["z_robusto"])

    subtitulo("Valores repetidos e redondos")
    rep = valores_repetidos(s, 5)
    if len(rep):
        tabela(["valor", "vezes", "% dos registros", "soma"],
               [[br(r["valor"], c), br(int(r["vezes"])), pct(r["pct"]), br(r["soma"], c)]
                for _, r in rep.iterrows()], ["r", "r", "r", "r"], recuo=4)
    red = valores_redondos(s)
    partes = []
    if red["n_1000"] >= 20:
        partes.append(f"{pct(red['mult_1000'])} dos valores ≥ 1.000 são múltiplos de 1.000")
    if red["n_100"] >= 20:
        partes.append(f"{pct(red['mult_100'])} dos ≥ 100 são múltiplos de 100")
    if partes:
        nota("; ".join(partes) + ". Valores redondos demais sugerem estimativa ou valor combinado.", 4)

    if cd or ce:
        subtitulo("Possível pagamento em duplicidade")
        chave = " + ".join([cv] + ([str(ce)] if ce else []) + ([f"{cd} (dia)" if ce else f"{cd} (minuto)"] if cd else []))
        dup = possiveis_duplicatas(df, cv, cd, ce)
        if dup.empty:
            _status("ok", f"Nenhum grupo repetido pela chave {chave}.", 4)
        else:
            grupos = dup.groupby("grupo_dup")
            excesso = float((grupos[cv].first() * (grupos.size() - 1)).sum())
            aviso(f"{br(dup['grupo_dup'].nunique())} grupos ({br(len(dup))} linhas) com a mesma chave "
                  f"({chave}); {br(excesso, c)} pagos a mais se forem duplicidade.", 4)
            _mostrar_linhas(rel, dup.head(rel.top * 2), max_linhas=rel.top * 2, extras=["grupo_dup"])
            rel.tabelas["Duplicidade"] = rel.com_linha(dup)
            rel.achado(f"Possível duplicidade: {br(dup['grupo_dup'].nunique())} grupos com mesma "
                       f"chave ({chave}), {br(excesso, c)} em excesso.", "alerta")

    subtitulo("Fracionamento — acúmulo logo abaixo de limites")
    limites = [rel.limite] if rel.limite else None
    acum = acumulo_abaixo_de_limites(s, limites)
    suspeitos = acum[acum["suspeito"]]
    if rel.limite:
        r = acum.iloc[0]
        nota(f"Limite informado: {br(rel.limite, 0)}. Entre {br(rel.limite * 0.95, 0)} e "
             f"{br(rel.limite, 0)}: {br(int(r['logo_abaixo']))} valores "
             f"(esperado ≈ {br(r['esperado'], 1)}; {br(r['razao'], 1)}×).", 4)
    if suspeitos.empty:
        _status("ok", "Nenhum acúmulo suspeito logo abaixo de valores redondos.", 4)
    for _, r in suspeitos.iterrows():
        L = r["limite"]
        aviso(f"Acúmulo logo abaixo de {br(L, 0)}: {br(int(r['logo_abaixo']))} valores entre "
              f"{br(L * 0.95, 0)} e {br(L, 0)}, {br(r['razao'], 1)}× o esperado. Pode ser "
              "fracionamento para fugir de alçada (ou preço psicológico, tipo 9,90).", 4)
        rel.achado(f"Acúmulo de valores logo abaixo de {br(L, 0)} ({br(r['razao'], 1)}× o esperado): "
                   "possível fracionamento.", "alerta")
    lim_fr = [rel.limite] if rel.limite else list(suspeitos["limite"])
    if ce and lim_fr:
        for L in lim_fr[:2]:
            fr = fracionamento(df, cv, L, ce, cd)
            if len(fr):
                aviso(f"{br(len(fr))} casos de {ce}" + (" + dia" if cd else "") + f" com 2+ valores "
                      f"abaixo de {br(L, 0)} somando acima dele:", 4)
                mostrar(fr.head(rel.top * 2), casas=c, recuo=6, max_linhas=rel.top * 2)
                rel.tabelas[f"Fracionamento {br(L, 0)}"] = fr.reset_index()
                rel.achado(f"{br(len(fr))} casos de possível fracionamento abaixo de {br(L, 0)} "
                           f"(mesmo {ce}" + (", mesmo dia)." if cd else ")."), "alerta")
    elif not ce:
        nota("Com -e (entidade) e --limite, o toolkit também procura o mesmo cliente/fornecedor "
             "quebrando um valor grande em vários pequenos no mesmo dia.", 4)

    if cd:
        subtitulo("Horário e calendário")
        cal = marcar_calendario(df[cd], rel.horas)
        validas = df[cd].notna()
        total_v = s[validas].sum()
        linhas = []
        for rot, mask in (("fim de semana", cal["fim_de_semana"]), ("feriado nacional", cal["feriado"].notna()),
                          ("madrugada (0h–5h)", cal.get("madrugada"))):
            if mask is None:
                continue
            q = int(mask.sum())
            linhas.append([rot, br(q), pct(q / max(validas.sum(), 1) * 100), br(s[mask].sum(), c),
                           pct(s[mask].sum() / total_v * 100) if total_v else "-"])
        tabela(["quando", "qtd", "% qtd", "total", "% total"], linhas, ["l", "r", "r", "r", "r"], recuo=4)
        fer = cal["feriado"].dropna()
        if len(fer):
            nota("Feriados com movimento: " + ", ".join(f"{n} ({q})" for n, q in fer.value_counts().head(4).items()), 4)
        if "madrugada" in cal and cal["madrugada"].mean() >= 0.02:
            rel.achado(f"{pct(cal['madrugada'].mean() * 100)} dos registros acontecem de madrugada (0h–5h).",
                       "aviso")

    subtitulo("Registros para conferir primeiro")
    al = sinalizar(df, cv, cd, ce, rel.limite if rel.limite else None, rel.horas, log=usa_log)
    al = al[al["pontos_risco"] > 0]
    al = al.loc[al.sort_values(["pontos_risco", cv], ascending=False).index]
    rel.alertas = al
    if al.empty:
        _status("ok", "Nenhum registro sinalizado.", 4)
        return
    pesos = ", ".join(f"{_NOMES_ALERTA[k]} {v}" for k, v in _PESOS_ALERTA.items())
    nota(f"pts = soma dos alertas de cada registro ({pesos}; fds = fim de semana).", 4)
    vista = al.drop(columns=[x for x in al.columns if str(x).startswith("alerta_")])
    _mostrar_linhas(rel, vista.rename(columns={"pontos_risco": "pts"}), max_linhas=15,
                    extras=["pts", "motivos"], max_cel=40)
    nota(f"{br(len(al))} registros com algum alerta; {br(int((al['pontos_risco'] >= 3).sum()))} "
         "com 3+ pontos. Lista completa: rel.alertas (python -i) ou --exportar.", 4)
    exportar = rel.com_linha(al)
    frente = [exportar.columns[0], "pontos_risco", "motivos"]
    rel.tabelas["Alertas"] = exportar[frente + [c for c in exportar.columns if c not in frente]]
    if (al["pontos_risco"] >= 3).any():
        rel.achado(f"{br(int((al['pontos_risco'] >= 3).sum()))} registros acumulam 3+ pontos de "
                   "risco — comece a conferência por eles.", "alerta")


def _secao_achados(rel):
    if not rel.achados:
        return
    rel.secao("PRINCIPAIS ACHADOS")
    ordem = {"alerta": 0, "aviso": 1, "info": 2}
    vistos, itens = set(), []
    for nivel, texto in sorted(rel.achados, key=lambda a: ordem.get(a[0], 3)):
        if texto not in vistos:
            vistos.add(texto)
            itens.append((nivel, texto))
    print()
    for i, (nivel, texto) in enumerate(itens, 1):
        marca = {"alerta": "!!", "aviso": "!", "info": "·"}.get(nivel, "·")
        estilo = {"alerta": "alerta", "aviso": "aviso", "info": "fraco"}.get(nivel, "fraco")
        _escrever(texto, 2, f"{i:>2}. " + pintar(f"{marca:<2}", estilo))
    rel.tabelas = {"Achados": pd.DataFrame(itens, columns=["nivel", "achado"]), **rel.tabelas}


def analisar(df, col_valor=None, col_categoria=None, col_data=None, top=5, col_entidade=None,
             fraude=False, limite=None, secoes=None, info=None, debug=False, _padroes=None):
    """Relatório completo no terminal. Devolve um Relatorio (rel.df, rel.tabelas,
    rel.achados, rel.alertas).

    col_*: colunas a usar (None = escolhe/sugere sozinho). fraude=True inclui a
    seção antifraude; limite = limite de alçada para o teste de fracionamento.
    secoes: lista com o que mostrar (padrão: tudo) — veja SECOES."""
    rel = Relatorio(df, info, top, debug)
    rel.limite = limite
    _escolher_colunas(rel, col_valor, col_categoria, col_data, col_entidade, _padroes)
    pedidas = set(secoes) if secoes else set(SECOES) - ({"fraude"} if not fraude else set())
    desconhecidas = pedidas - set(SECOES)
    if desconhecidas:
        raise ValueError(f"Seções desconhecidas: {', '.join(sorted(desconhecidas))}. "
                         f"Use: {', '.join(SECOES)}")
    plano = [
        ("visao", _secao_visao, True),
        ("qualidade", _secao_qualidade, True),
        ("ficha", _secao_ficha, rel.col_valor is not None),
        ("outliers", _secao_outliers, rel.col_valor is not None),
        ("extremos", _secao_extremos, rel.col_valor is not None),
        ("categorias", _secao_categorias, bool(rel.col_categoria or rel.cand_categoria)),
        ("entidades", _secao_entidades, rel.col_entidade is not None),
        ("tempo", _secao_tempo, rel.col_data is not None),
        ("correlacao", _secao_correlacao, True),
        ("fraude", _secao_fraude, rel.col_valor is not None),
    ]
    for nome, funcao, aplicavel in plano:
        if nome in pedidas and aplicavel:
            _rodar(rel, nome, funcao)
        elif nome == "fraude" and nome in pedidas and rel.col_valor is None:
            print()
            aviso("Antifraude precisa de uma coluna numérica de valor (-v).")
    if "achados" in pedidas:
        _rodar(rel, "achados", _secao_achados)
    return rel


# ===========================================================================
# 11. EXPORTAÇÃO
# ===========================================================================
def _nome_aba(nome, usados):
    base = re.sub(r"[\[\]:*?/\\]", "-", str(nome))[:31] or "Aba"
    nome, i = base, 2
    while nome.lower() in usados:
        nome = f"{base[:28]}_{i}"
        i += 1
    usados.add(nome.lower())
    return nome


def exportar_excel(rel, caminho):
    """Salva os achados e as tabelas do relatório num .xlsx (uma aba por tabela),
    com cabeçalho fixo, filtro e números formatados."""
    caminho = Path(caminho)
    if caminho.suffix.lower() != ".xlsx":
        caminho = caminho.with_suffix(".xlsx")
    tabelas = dict(rel.tabelas)
    if "Achados" not in tabelas and rel.achados:
        tabelas = {"Achados": pd.DataFrame(rel.achados, columns=["nivel", "achado"]), **tabelas}
    if not tabelas:
        raise ErroToolkit("Nada para exportar.")
    try:
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as erro:
        raise ErroToolkit("Para exportar instale: pip install openpyxl") from erro
    usados = set()
    try:
        with pd.ExcelWriter(caminho, engine="openpyxl") as w:
            for nome, t in tabelas.items():
                t = t.copy()
                for col in t.columns:
                    if pd.api.types.is_datetime64_any_dtype(t[col]):
                        t[col] = _sem_fuso(t[col])
                    elif t[col].map(lambda v: isinstance(v, (list, dict))).any():
                        t[col] = t[col].map(lambda v: ", ".join(map(str, v)) if isinstance(v, list) else v)
                aba = _nome_aba(nome, usados)
                t.to_excel(w, sheet_name=aba, index=False)
                ws = w.sheets[aba]
                ws.freeze_panes = "A2"
                if len(t.columns):
                    ws.auto_filter.ref = ws.dimensions
                for i, col in enumerate(t.columns, 1):
                    cel = ws.cell(row=1, column=i)
                    cel.font = Font(bold=True, color="FFFFFF")
                    cel.fill = PatternFill("solid", fgColor="305496")
                    amostra = t[col].head(300).map(lambda v: len(_celula(v, max_cel=60)))
                    ws.column_dimensions[get_column_letter(i)].width = min(
                        60, max([len(str(col))] + amostra.tolist()) + 2)
                    if pd.api.types.is_float_dtype(t[col]):
                        fmt = "#,##0.00"
                    elif pd.api.types.is_integer_dtype(t[col]) and not _inteiro_sem_milhar(t[col]):
                        fmt = "#,##0"
                    elif pd.api.types.is_datetime64_any_dtype(t[col]):
                        fmt = "dd/mm/yyyy hh:mm"
                    else:
                        continue
                    for linha in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                        linha[0].number_format = fmt
    except PermissionError as erro:
        raise ErroToolkit(f"Não consegui salvar {caminho.name}: ele está aberto no Excel? "
                          "Feche e rode de novo.") from erro
    return caminho


class _Espelho(io.TextIOBase):
    """Escreve no terminal e num arquivo ao mesmo tempo (sem as cores)."""

    def __init__(self, caminho):
        self.arquivo = open(caminho, "w", encoding="utf-8")
        self.original = sys.stdout
        sys.stdout = self

    def write(self, texto):
        self.original.write(texto)
        self.arquivo.write(_RE_ANSI.sub("", texto))
        return len(texto)

    def flush(self):
        self.original.flush()
        if not self.arquivo.closed:
            self.arquivo.flush()

    def fechar(self):
        sys.stdout = self.original
        self.arquivo.close()


# ===========================================================================
# 12. DADOS DE EXEMPLO
# ===========================================================================
def _gerar_cpf(rng, valido=True):
    d = list(rng.integers(0, 10, 9))
    for pesos in (range(10, 1, -1), range(11, 1, -1)):
        d.append(sum(a * b for a, b in zip(d, pesos)) * 10 % 11 % 10)
    if not valido:
        d[-1] = (d[-1] + 1) % 10
    s = "".join(map(str, d))
    return f"{s[:3]}.{s[3:6]}.{s[6:9]}-{s[9:]}"


def dados_exemplo(n=3000, seed=7, como_texto=True):
    """Base fictícia de transações com problemas plantados (duplicatas,
    outliers, fracionamento abaixo de 10.000, madrugada, CPF inválido, grafias
    diferentes, datas futuras...). como_texto=True devolve Valor e Data como
    texto BR ('R$ 1.234,56', '31/01/2026 14:30'), igual a um CSV exportado:
    use preparar(dados_exemplo()) ou python toolkit_pandas.py --demo."""
    rng = np.random.default_rng(seed)
    nomes = ["Ana", "Bruno", "Carla", "Diego", "Elisa", "Fábio", "Gabriela", "Heitor", "Isabela",
             "João", "Larissa", "Marcos", "Natália", "Otávio", "Paula", "Rafael", "Sabrina",
             "Tiago", "Vanessa", "Wagner"]
    sobrenomes = ["Silva", "Souza", "Oliveira", "Santos", "Lima", "Pereira", "Costa", "Almeida",
                  "Ferreira", "Rodrigues", "Gomes", "Martins", "Araújo", "Barbosa", "Ribeiro"]
    clientes = list(dict.fromkeys(f"{a} {b}" for a in nomes for b in sobrenomes))
    rng.shuffle(clientes)
    clientes = clientes[:160]
    cpfs = [_gerar_cpf(rng, valido=i not in (17, 58, 101)) for i in range(len(clientes))]
    peso = 1 / np.arange(1, len(clientes) + 1) ** 0.8
    cli = rng.choice(len(clientes), n, p=peso / peso.sum())

    tipos = rng.choice(["PIX", "Boleto", "Cartão", "TED", "DOC"], n, p=[.45, .2, .18, .13, .04])
    mult = pd.Series({"PIX": 1, "Boleto": 1.6, "Cartão": .55, "TED": 3.2, "DOC": 2})
    valor = rng.lognormal(np.log(300), 1.05, n) * mult[tipos].to_numpy()

    inicio = pd.Timestamp("2025-10-01")
    dias = rng.integers(0, 365, n * 3)
    dow = (inicio + pd.to_timedelta(dias, "D")).dayofweek
    aceita = rng.random(n * 3) < np.array([1, 1, 1, 1, 1, .45, .3])[dow]
    dias = dias[aceita][:n]
    horas = np.clip(rng.normal(14, 3.2, n), 7, 21.99)
    madrugada = rng.random(n) < 0.025
    horas[madrugada] = rng.uniform(0, 5.99, madrugada.sum())
    datas = inicio + pd.to_timedelta(dias, "D") + pd.to_timedelta((horas * 60).astype(int), "min")

    df = pd.DataFrame({
        "Data": datas, "Cliente": [clientes[i] for i in cli], "CPF": [cpfs[i] for i in cli],
        "Tipo": tipos, "Canal": rng.choice(["App", "Internet Banking", "Agência"], n, p=[.55, .3, .15]),
        "UF": rng.choice(["SP", "RJ", "MG", "PR", "RS", "BA", "SC", "PE", "GO", "DF"], n,
                         p=[.32, .14, .12, .08, .07, .07, .06, .05, .05, .04]),
        "Status": rng.choice(["Aprovada", "Recusada", "Pendente"], n, p=[.9, .07, .03]),
        "Valor": valor.round(2),
    })
    i = rng.choice(n, 20, replace=False)                                  # outliers
    df.loc[i, "Valor"] = (df.loc[i, "Valor"] * rng.uniform(25, 60, 20)).round(2)
    i = rng.choice(n, int(n * 0.04), replace=False)                       # redondos
    df.loc[i, "Valor"] = rng.choice([500.0, 1000.0, 1500.0, 2000.0, 5000.0], len(i))
    i = rng.choice(n, int(n * 0.015), replace=False)                      # estornos
    df.loc[i, "Valor"] = -df.loc[i, "Valor"].abs()

    extras = []                                                           # fracionamento
    for k in (9, 23, 41):
        for dia in rng.choice(365, 12, replace=False):
            for _ in range(rng.integers(2, 5)):
                extras.append({"Data": inicio + pd.Timedelta(days=int(dia), hours=int(rng.integers(9, 18)),
                                                             minutes=int(rng.integers(0, 60))),
                               "Cliente": clientes[k], "CPF": cpfs[k], "Tipo": "TED", "Canal": "Internet Banking",
                               "UF": "SP", "Status": "Aprovada", "Valor": round(rng.uniform(9500, 9990), 2)})
    base_dup = df.sample(12, random_state=seed).copy()                   # duplicidade
    base_dup["Data"] = base_dup["Data"] + pd.to_timedelta(rng.integers(1, 40, 12), "min")
    df = pd.concat([df, pd.DataFrame(extras), base_dup], ignore_index=True)

    m = len(df)
    variantes = {"PIX": ["pix", "Pix "], "Cartão": ["Cartao", "CARTÃO"], "Boleto": ["boleto"]}
    sorteio = pd.Series(False, index=df.index)
    sorteio.iloc[rng.choice(m, int(m * 0.025), replace=False)] = True
    for certo, erradas in variantes.items():
        alvo = sorteio & (df["Tipo"] == certo)
        df.loc[alvo, "Tipo"] = rng.choice(erradas, int(alvo.sum()))
    df.loc[rng.choice(m, 25, replace=False), "UF"] = rng.choice(["sp", "S.P.", "São Paulo"], 25)
    df.loc[rng.choice(m, int(m * 0.02), replace=False), "UF"] = None
    df.loc[rng.choice(m, int(m * 0.01), replace=False), "Canal"] = None
    df.loc[rng.choice(m, 2, replace=False), "Data"] += pd.DateOffset(years=1)  # digitação errada

    df = df.sort_values("Data", kind="stable").reset_index(drop=True)
    df.insert(0, "ID_Transacao", [f"TX{i:06d}" for i in range(1, len(df) + 1)])
    df = pd.concat([df, df.sample(4, random_state=seed + 1)], ignore_index=True)  # cópias exatas
    df.loc[rng.choice(len(df), 8, replace=False), "Valor"] = np.nan

    if como_texto:
        df["Valor"] = df["Valor"].astype(object)
        def rs(v):
            if _na(v):
                return None
            t = f"{abs(v):,.2f}".replace(",", "\0").replace(".", ",").replace("\0", ".")
            return f"-R$ {t}" if v < 0 else f"R$ {t}"
        df["Valor"] = df["Valor"].map(rs)
        df.loc[rng.choice(len(df), 2, replace=False), "Valor"] = "isento"
        df["Data"] = df["Data"].dt.strftime("%d/%m/%Y %H:%M")
    return df


# ===========================================================================
# 13. AJUDA E COLA
# ===========================================================================
_AJUDA = [
    ("Carga", [
        ("carregar(caminho, aba=, cabecalho=)", "lê csv/xlsx/json/parquet arrumando números BR e datas"),
        ("preparar(df)", "mesma limpeza para um df vindo de outro lugar (SQL, clipboard)"),
        ("converter_numero(serie, 'br')", "'R$ 1.234,56' -> 1234.56"),
        ("converter_data(serie)", "'31/01/2026' -> data (dia/mês por padrão)"),
    ]),
    ("Ver", [
        ("mostrar(df, max_linhas=20)", "imprime alinhado, formato BR, cabe no terminal"),
        ("visao_geral(df) / dicionario_dados(df)", "tipo, papel, nulos e exemplo por coluna"),
        ("br(x) / pct(x)", "1234.5 -> '1.234,50' / 12.3 -> '12,3%'"),
    ]),
    ("Qualidade", [
        ("checar_qualidade(df)", "regras: datas futuras, IDs repetidos, grafias, CPF inválido..."),
        ("diagnostico_nulos(df) / duplicatas(df)", "vazios por coluna / linhas repetidas"),
        ("tratar_nulos(df, 'mediana')", "cópia com nulos tratados (ou removidos)"),
        ("grafias_inconsistentes(serie)", "'SP', 'sp ', 'S.P.' no mesmo grupo"),
        ("padronizar_texto(serie) / padronizar_uf(serie)", "limpa texto / 'São Paulo' -> 'SP'"),
        ("detectar_pii(df) / mascarar_pii(df)", "acha e mascara CPF, e-mail, telefone, nome"),
        ("validar_cpf(x) / validar_cnpj(x)", "dígitos verificadores (CNPJ alfanumérico incluso)"),
    ]),
    ("Estatística", [
        ("resumo_numerico(serie)", "média, mediana, moda, desvio, CV, quartis..."),
        ("leitura_em_frase(serie)", "os números traduzidos em texto"),
        ("histograma(serie)", "distribuição em faixas"),
        ("outliers_iqr(df, col) / outliers_mad(df, col)", "fora do padrão (IQR / z robusto)"),
        ("correlacao(df) / pares_correlacao(df)", "matriz / pares mais relacionados"),
    ]),
    ("Agrupar", [
        ("agrupar(df, chave, valor)", "groupby com soma, média e contagem"),
        ("cruzar(df, linhas, colunas, valor)", "tabela dinâmica com totais"),
        ("curva_abc(df, chave, valor)", "Pareto: quem faz 80% do total"),
        ("frequencia(df, coluna)", "contagem com % e % acumulado"),
        ("serie_temporal(df, data, valor, 'MS')", "soma por dia/semana/mês/trimestre"),
        ("resumo_temporal(df, data, valor)", "por período com variação e períodos parciais"),
        ("coorte(df, entidade, data)", "retenção por mês de entrada"),
    ]),
    ("Antifraude", [
        ("benford(serie, digitos=1)", "Lei de Benford com MAD e conformidade (Nigrini)"),
        ("zscore_robusto(serie)", "z-score modificado (mediana/MAD)"),
        ("possiveis_duplicatas(df, valor, data, entidade)", "mesmo valor, mesma entidade, mesmo dia"),
        ("acumulo_abaixo_de_limites(serie)", "degraus logo abaixo de 1.000, 5.000, 10.000..."),
        ("fracionamento(df, valor, limite, entidade, data)", "valor grande quebrado em pequenos"),
        ("valores_repetidos(serie) / valores_redondos(serie)", "duplicação de números / redondos"),
        ("marcar_calendario(datas) / feriados_br(ano)", "fim de semana, feriado, madrugada"),
        ("sinalizar(df, valor, data, entidade)", "todos os alertas + pontos_risco por linha"),
    ]),
    ("Juntar tabelas", [
        ("checar_merge(esq, dir, chave)", "ANTES do merge: casamento, repetidas, tipos, explosão"),
        ("padronizar_chave(serie)", "chave comparável ('00123 ' == 123.0)"),
    ]),
    ("Relatório", [
        ("analisar(df, col_valor, ...)", "relatório completo; devolve rel"),
        ("exportar_excel(rel, 'saida.xlsx')", "achados e tabelas em abas do Excel"),
        ("dados_exemplo()", "base fictícia com problemas plantados"),
        ("cola()", "filtros e comandos de pandas que você sempre esquece"),
    ]),
]


def ajuda(funcao=None):
    """Lista as funções do toolkit. ajuda(mostrar) mostra os detalhes de uma."""
    if funcao is not None:
        nome = getattr(funcao, "__name__", str(funcao))
        print(f"\n{nome}{_assinatura(funcao)}\n")
        print(textwrap.indent(textwrap.dedent(funcao.__doc__ or "(sem documentação)").strip(), "  "))
        return
    for grupo, itens in _AJUDA:
        subtitulo(grupo)
        larg = max(len(a) for a, _ in itens)
        sobra = max(20, _SAIDA["largura"] - larg - 6)
        for assinatura, descricao in itens:
            partes = textwrap.wrap(descricao, sobra) or [""]
            print(f"    {assinatura:<{larg}}  {pintar(partes[0], 'fraco')}")
            for resto in partes[1:]:
                print(f"    {'':<{larg}}  {pintar(resto, 'fraco')}")
    print()
    nota("Detalhes de uma função: ajuda(nome_da_funcao)  ·  variáveis: df, rel (rel.alertas, rel.tabelas)")


def _assinatura(funcao):
    try:
        import inspect
        return str(inspect.signature(funcao))
    except (TypeError, ValueError):
        return ""


_COLA = """\
SELECIONAR
  df['Col'] / df[['A', 'B']]               coluna / colunas
  df.loc[10, 'Valor']                      por RÓTULO (linha, coluna)
  df.iloc[0, 2]                            por POSIÇÃO
FILTRAR
  df[df['Valor'] > 100]                    simples
  df[(df['A'] > 0) & (df['B'] == 'x')]     E   (parênteses obrigatórios)
  df[(df['A'] > 0) | (df['B'] == 'x')]     OU
  df[~df['Tipo'].isin(['PIX', 'TED'])]     NÃO está na lista
  df[df['Valor'].between(100, 500)]        intervalo (inclusivo)
  df[df['Nome'].str.contains('jo', case=False, na=False)]   texto
  df[df['Data'] >= '2026-01-01']           data
  df.query("Valor > 100 and Tipo == 'PIX'")    legível (`Nome Com Espaço` entre crases)
CRIAR / MUDAR
  df['Taxa'] = df['Valor'] * 0.05          coluna calculada
  df['Faixa'] = pd.cut(df['Valor'], [0, 100, 1000, float('inf')])
  df['Mes'] = df['Data'].dt.to_period('M') mês da data
  df.rename(columns={'old': 'new'}) / df.drop(columns=['X'])
AGRUPAR
  df.groupby('Tipo')['Valor'].agg(['sum', 'mean', 'count'])
  df.groupby(['Tipo', 'UF'], as_index=False)['Valor'].sum()
  df.pivot_table(index='UF', columns='Tipo', values='Valor', aggfunc='sum', margins=True)
  df['Tipo'].value_counts(normalize=True)  proporção de cada valor
JUNTAR
  checar_merge(a, b, 'CPF')                SEMPRE antes do merge
  a.merge(b, on='CPF', how='left', validate='m:1', indicator=True)
  pd.concat([a, b], ignore_index=True)     empilhar
ORDENAR / TOP
  df.sort_values('Valor', ascending=False) / df.nlargest(10, 'Valor')
SALVAR
  df.to_excel('saida.xlsx', index=False) / df.to_csv('saida.csv', sep=';', decimal=',', index=False, encoding='utf-8-sig')
"""


def cola():
    """Filtros e comandos de pandas que você sempre esquece."""
    print()
    for linha in _COLA.splitlines():
        if linha and not linha.startswith(" "):
            print(pintar(linha, "titulo"))
        else:
            print(linha)


# ===========================================================================
# 14. LINHA DE COMANDO
# ===========================================================================
def _numero_cli(texto):
    t = str(texto).strip().lower().replace(" ", "")
    fator = 1
    for sufixo, f in (("mil", 1e3), ("k", 1e3), ("mi", 1e6), ("m", 1e6)):
        if t.endswith(sufixo):
            t, fator = t[: -len(sufixo)], f
            break
    for formato in ("br", "us"):
        v = converter_numero(pd.Series([t], dtype=object), formato).iloc[0]
        if not _na(v):
            return v * fator
    raise argparse.ArgumentTypeError(f"número inválido: {texto!r} (ex: 10000, 10.000, 10k)")


def _aplicar_filtro(df, expressao):
    try:
        out = df.query(expressao)
    except Exception:
        try:
            out = df.query(expressao, engine="python")
        except Exception as erro:
            raise ErroToolkit(
                f"Filtro inválido: {expressao!r}\n{type(erro).__name__}: {erro}\n"
                "Exemplos: \"Valor > 1000\"  \"Tipo == 'PIX'\"  \"Data >= '2026-01-01'\"  "
                "\"`Valor Total` > 0\" (nome com espaço entre crases)") from erro
    _status("info", f"Filtro {expressao!r}: {br(len(out))} de {br(len(df))} linhas "
                    f"({pct(len(out) / max(len(df), 1) * 100)}).")
    if out.empty:
        raise ErroToolkit("O filtro não deixou nenhuma linha.")
    return out


def _banner(info, df):
    largura = _SAIDA["largura"]
    nome = info.get("arquivo", "DataFrame")
    marca = f"toolkit_pandas {__version__}"
    print(pintar("═" * largura, "titulo"))
    espaco = max(2, largura - len(nome) - len(marca) - 4)
    print(pintar(f"  {nome}{' ' * espaco}{marca}", "titulo"))
    print(pintar("═" * largura, "titulo"))
    _imprimir_carga(info, df)


def _mensagem_erro(erro):
    if isinstance(erro, UnicodeDecodeError):
        return ("Não consegui ler o texto do arquivo (encoding). Tente --encoding utf-8 "
                "ou --encoding latin-1.")
    if isinstance(erro, PermissionError):
        return f"Sem permissão para acessar {erro.filename or 'o arquivo'}. Ele está aberto no Excel?"
    if isinstance(erro, KeyError) and not isinstance(erro, ColunaNaoEncontrada):
        return f"Coluna não encontrada: {erro}"
    return str(erro.args[0]) if getattr(erro, "args", None) else str(erro)


def main(argv=None):
    p = argparse.ArgumentParser(
        prog="toolkit_pandas.py",
        description="Análise exploratória no terminal: qualidade, estatística, categorias, "
                    "tempo e antifraude, com números no formato brasileiro.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        add_help=False,
        epilog=textwrap.dedent("""\
            exemplos:
              python toolkit_pandas.py vendas.xlsx
              python toolkit_pandas.py vendas.xlsx -v Total -c Loja -d Data
              python toolkit_pandas.py extrato.csv -f -e Cliente --limite 10000
              python toolkit_pandas.py base.csv -q "Status == 'Aprovada'" --exportar achados.xlsx
              python toolkit_pandas.py base.csv -s visao,qualidade      (só olhar a base)
              python toolkit_pandas.py --demo                           (dados de exemplo)
              python -i toolkit_pandas.py base.csv                      (explorar depois: ajuda())
        """),
    )
    p.add_argument("arquivo", nargs="?", help="csv, txt, tsv, xlsx, xls, xlsm, json, jsonl ou parquet")
    g = p.add_argument_group("colunas (nome com espaço vai entre aspas)")
    g.add_argument("-v", "--valor", help="coluna numérica principal (padrão: escolhe sozinho)")
    g.add_argument("-c", "--categoria", help="coluna para agrupar")
    g.add_argument("-d", "--data", help="coluna de data (padrão: a única que houver)")
    g.add_argument("-e", "--entidade", help="quem transaciona: cliente, fornecedor, conta...")
    g = p.add_argument_group("análise")
    g.add_argument("-f", "--fraude", action="store_true", help="inclui a seção antifraude")
    g.add_argument("--limite", type=_numero_cli, help="limite de alçada p/ fracionamento (ex: 10000)")
    g.add_argument("-q", "--filtro", help="filtra antes de analisar (sintaxe do df.query)")
    g.add_argument("-n", "--top", type=int, default=TOP, help=f"linhas nas listas (padrão {TOP})")
    g.add_argument("-s", "--secoes", help=f"só estas seções, separadas por vírgula: {','.join(SECOES)}")
    g = p.add_argument_group("leitura")
    g.add_argument("--aba", help="aba do Excel (nome ou número, 1 = primeira)")
    g.add_argument("--cabecalho", type=int, help="linha do cabeçalho (1 = primeira; padrão: detecta)")
    g.add_argument("--sep", help="separador do CSV (padrão: detecta)")
    g.add_argument("--encoding", help="encoding do CSV (padrão: detecta)")
    g = p.add_argument_group("saída")
    g.add_argument("--exportar", metavar="ARQ.xlsx", help="salva achados e tabelas em Excel")
    g.add_argument("--salvar", metavar="ARQ.txt", help="salva o relatório em texto")
    g.add_argument("--sem-cor", action="store_true", help="desliga as cores")
    g.add_argument("--largura", type=int, help="largura do relatório em caracteres")
    g.add_argument("--demo", action="store_true", help="roda com dados fictícios de exemplo")
    g.add_argument("--debug", action="store_true", help="mostra o traceback dos erros")
    g.add_argument("-h", "--ajuda", "--help", action="help", help="mostra esta ajuda")
    g.add_argument("--versao", action="version", version=f"toolkit_pandas {__version__}")
    args = p.parse_args(argv)

    configurar_saida(cor=False if args.sem_cor else None, largura=args.largura)
    espelho = _Espelho(args.salvar) if args.salvar else None
    try:
        if args.demo:
            info = {"arquivo": "dados de exemplo (--demo)", "formato": "DataFrame",
                    "mensagens": [], "linha_offset": None}
            df = preparar(dados_exemplo(), verbose=False, _info=info)
            args.categoria = args.categoria or "Tipo"
            args.entidade = args.entidade or "Cliente"
            args.fraude = True
        else:
            arquivo = args.arquivo or ARQUIVO
            if arquivo is None:
                p.print_usage()
                print("\nInforme o arquivo a analisar (ou --demo para ver um exemplo)."
                      + _sugerir_arquivos(Path.cwd()))
                return None
            kwargs = dict(LEITURA_KWARGS)
            if args.sep:
                kwargs["sep"] = args.sep.replace("\\t", "\t")
            if args.encoding:
                kwargs["encoding"] = args.encoding
            df, info = _carregar(arquivo, aba=args.aba, cabecalho=args.cabecalho, **kwargs)
        _banner(info, df)
        if args.filtro:
            df = _aplicar_filtro(df, args.filtro)
        secoes = [s.strip() for s in args.secoes.split(",")] if args.secoes else None
        if secoes and "fraude" in secoes:
            args.fraude = True
        rel = analisar(
            df, col_valor=args.valor, col_categoria=args.categoria, col_data=args.data,
            top=args.top, col_entidade=args.entidade, fraude=args.fraude or FRAUDE,
            limite=args.limite or LIMITE_ALCADA, secoes=secoes, info=info, debug=args.debug,
            _padroes={"valor": COL_VALOR, "categoria": COL_CATEGORIA, "data": COL_DATA,
                      "entidade": COL_ENTIDADE},
        )
        print()
        print(pintar("═" * _SAIDA["largura"], "titulo"))
        if args.exportar:
            destino = exportar_excel(rel, args.exportar)
            _status("ok", f"Tabelas exportadas para {destino} ({len(rel.tabelas)} abas).")
        if args.salvar:
            _status("ok", f"Relatório salvo em {args.salvar}.")
        if not sys.flags.interactive:
            nota("Para explorar: python -i toolkit_pandas.py ... → df, rel e ajuda() ficam disponíveis.")
        else:
            nota("df e rel estão na memória. Digite ajuda() para ver as funções.")
        return rel
    except (ErroToolkit, FileNotFoundError, KeyError, ValueError, OSError, ImportError,
            pd.errors.ParserError) as erro:
        print(f"\nERRO: {_mensagem_erro(erro)}", file=sys.stderr)
        if args.debug:
            traceback.print_exc()
        return None
    except Exception as erro:
        print(f"\nERRO inesperado: {type(erro).__name__}: {erro}\n"
              "Rode de novo com --debug para ver os detalhes.", file=sys.stderr)
        if args.debug:
            traceback.print_exc()
        return None
    finally:
        if espelho:
            espelho.fechar()


if __name__ == "__main__":
    for _fluxo in (sys.stdout, sys.stderr):                 # acentos e barras no Windows
        if hasattr(_fluxo, "reconfigure"):
            _fluxo.reconfigure(encoding="utf-8", errors="replace")
    rel = main()
    df = rel.df if rel is not None else None
    if rel is None and not sys.flags.interactive:
        sys.exit(1)
